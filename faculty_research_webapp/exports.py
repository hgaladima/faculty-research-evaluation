"""Build the downloadable evaluation report as an Excel workbook or a PDF."""
from __future__ import annotations

import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFError, TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INFO_LABELS = [("name", "Faculty member"), ("faculty_id", "Faculty ID"), ("rank", "Academic rank"),
               ("department", "Department"), ("appointment", "Appointment type"),
               ("evaluation_year", "Evaluation year"), ("research_effort", "Assigned research effort (%)"),
               ("months_active", "Months active")]
DISCLAIMER = ("Decision aid only. The scoring policy is provisional until adopted by the annual review "
              "committee; the chair's written evaluation and the faculty narrative remain the record.")

CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
FORMULA_STARTS = ("=", "+", "-", "@")
# A table cell taller than a page cannot be laid out, so long entries are shortened in the PDF tables only.
PDF_CELL_LIMIT = 600
SYSTEM_FONTS = [("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
                ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
                ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")]


def clean(value) -> str:
    """Text safe for both report formats: no None, no control characters."""
    return "" if value is None else CONTROL_CHARS.sub("", str(value))


def policy_stamp(result: dict, meta: dict) -> str:
    if meta.get("locked"):
        status = f"locked, approved by {meta.get('approved_by') or 'not recorded'}"
    else:
        status = "not locked (provisional, not approved)"
    return (f"{result['policy']['name']} | version {meta.get('version', 1)} | "
            f"fingerprint {result['policy']['fingerprint']} | {status}")


def summary_rows(result: dict) -> list[list]:
    rows = [["Research domain", "Points earned", "Maximum points", "Minimum met", "Performance"]]
    for d in result["domains"].values():
        rows.append([d["name"], d["score"], d["max"], "Yes" if d["minimum_met"] else "No", d["status"]])
    rows.append(["Overall research score", result["total"], 100, "", result["class"]])
    return rows


def breakdown_rows(result: dict) -> list[list]:
    rows = [["Domain", "Item", "How it was credited", "Units", "Counts toward minimum"]]
    for line in result["breakdown"]:
        rows.append([result["domains"][line["domain"]]["name"], line["item"], line["detail"],
                     round(line["units"], 3), line["toward_minimum"] or ""])
    return rows


def conversion_rows(result: dict) -> list[list]:
    rows = [["Domain", "Units earned", "Benchmark units", "Units for full excellence", "Points at the minimum",
             "Excellence points", "Domain points", "Basis for the minimum"]]
    for d in result["domains"].values():
        rows.append([d["name"], round(d["units"], 3), round(d["benchmark_units"], 3), round(d["saturation"], 3),
                     round(d["anchor"], 2), round(d["excellence"], 2), d["score"], d["minimum_basis"] or "not met"])
    return rows


def excel_bytes(info: dict, result: dict, policy: dict, meta: dict, context: dict) -> bytes:
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")

    def put(ws, row, column, value, bold=False):
        """Write a cell. Entered text is always stored as text, never as a formula."""
        if value is None or isinstance(value, str):
            value = clean(value)
        cell = ws.cell(row=row, column=column, value=value)
        if isinstance(value, str) and value.startswith(FORMULA_STARTS):
            cell.data_type = "s"
        if bold:
            cell.font = Font(bold=True)
        return cell

    def table(ws, rows, start):
        for r, row in enumerate(rows, start):
            for c, value in enumerate(row, 1):
                cell = put(ws, r, c, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if r == start:
                    cell.font, cell.fill = head_font, head_fill
        return start + len(rows) + 1

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Faculty Research Evaluation Report"
    ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for key, label in INFO_LABELS:
        put(ws, r, 1, label, bold=True)
        put(ws, r, 2, info.get(key))
        r += 1
    put(ws, r + 1, 1, "Research evaluation", bold=True)
    put(ws, r + 1, 2, result["class"])
    r += 3
    if result["scored"]:
        r = table(ws, summary_rows(result), r)
    for title, items in (("Explanation", result["notes"]), ("For committee attention", result["flags"]),
                         ("Input errors", result["errors"])):
        if items:
            put(ws, r, 1, title, bold=True)
            for item in items:
                r += 1
                put(ws, r, 1, item)
            r += 2
    for label, text in context.items():
        if text:
            put(ws, r, 1, f"{label} (recorded, not scored)", bold=True)
            put(ws, r + 1, 1, text).alignment = Alignment(wrap_text=True, vertical="top")
            r += 3
    put(ws, r, 1, "Scoring policy", bold=True)
    put(ws, r + 1, 1, policy_stamp(result, meta))
    put(ws, r + 2, 1, DISCLAIMER)
    for col, width in zip("ABCDE", (34, 16, 16, 14, 28)):
        ws.column_dimensions[col].width = width

    if result["scored"]:
        ws2 = wb.create_sheet("Points breakdown")
        nxt = table(ws2, conversion_rows(result), 1)
        table(ws2, breakdown_rows(result), nxt + 1)
        for i, width in enumerate((24, 34, 52, 16, 20, 16, 14, 40), 1):
            ws2.column_dimensions[get_column_letter(i)].width = width

    ws3 = wb.create_sheet("Policy")
    rows = [["Setting", "Value"]]
    for key in sorted(policy):
        value = policy[key]
        if isinstance(value, dict):
            rows += [[f"{key}.{k}", v] for k, v in value.items()]
        else:
            rows.append([key, ", ".join(value) if isinstance(value, list) else value])
    table(ws3, rows, 1)
    ws3.column_dimensions["A"].width, ws3.column_dimensions["B"].width = 44, 40

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_fonts() -> tuple[str, str]:
    """A regular and bold font. Uses a system font with wide character coverage when one is installed."""
    if "ReportBody" in pdfmetrics.getRegisteredFontNames():
        return "ReportBody", "ReportBody-Bold"
    for regular, bold in SYSTEM_FONTS:
        if Path(regular).exists() and Path(bold).exists():
            try:
                pdfmetrics.registerFont(TTFont("ReportBody", regular))
                pdfmetrics.registerFont(TTFont("ReportBody-Bold", bold))
                pdfmetrics.registerFontFamily("ReportBody", normal="ReportBody", bold="ReportBody-Bold",
                                              italic="ReportBody", boldItalic="ReportBody-Bold")
            except (TTFError, OSError):
                continue
            return "ReportBody", "ReportBody-Bold"
    return "Helvetica", "Helvetica-Bold"


def pdf_bytes(info: dict, result: dict, meta: dict, context: dict) -> bytes:
    regular, bold = pdf_fonts()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName=regular, fontSize=9.5, leading=12.5)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.grey)
    cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=10)
    title = ParagraphStyle("title", parent=styles["Title"], fontName=bold)
    heading = ParagraphStyle("heading", parent=styles["Heading3"], fontName=bold)
    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                       ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                       ("FONTNAME", (0, 0), (-1, -1), regular),
                       ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("VALIGN", (0, 0), (-1, -1), "TOP")])

    def para(value, style=body) -> Paragraph:
        return Paragraph(escape(clean(value)).replace("\n", "<br/>"), style)

    def cells(rows: list[list]) -> list[list]:
        def short(v):
            text = clean(v)
            return text if len(text) <= PDF_CELL_LIMIT else text[:PDF_CELL_LIMIT].rstrip() + " [...]"
        return [[para(short(v), cell) for v in row] for row in rows]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch,
                            title="Faculty Research Evaluation Report")
    story = [Paragraph("Faculty Research Evaluation Report", title)]
    info_rows = [[Paragraph(f"<b>{label}</b>", body), para(info.get(key))] for key, label in INFO_LABELS]
    story += [Table(info_rows, colWidths=[2.3 * inch, 4.6 * inch]), Spacer(1, 10)]
    story.append(Paragraph(f"<b>Research evaluation: {escape(result['class'])}</b>"
                           + (f" &nbsp; Overall research score: {result['total']} / 100" if result["scored"] else ""),
                           heading))
    if result["scored"]:
        t = Table(summary_rows(result), colWidths=[2.3 * inch, 1.1 * inch, 1.2 * inch, 1.0 * inch, 1.5 * inch])
        t.setStyle(grid)
        story += [t, Spacer(1, 8)]
    for heading_text, items in (("Explanation of the outcome", result["notes"]),
                                ("For committee attention", result["flags"]), ("Input errors", result["errors"])):
        if items:
            story.append(Paragraph(f"<b>{heading_text}</b>", body))
            story += [Paragraph(f"&bull; {escape(clean(item))}", body) for item in items]
            story.append(Spacer(1, 6))
    if result["scored"]:
        story.append(Paragraph("<b>How the points were earned</b>", body))
        t = Table(cells(conversion_rows(result)),
                  colWidths=[1.1 * inch, 0.6 * inch, 0.75 * inch, 0.8 * inch, 0.7 * inch, 0.75 * inch, 0.6 * inch,
                             1.8 * inch], repeatRows=1)
        t.setStyle(grid)
        story += [t, Spacer(1, 6)]
        t = Table(cells([row[:4] for row in breakdown_rows(result)]),
                  colWidths=[1.25 * inch, 1.9 * inch, 3.35 * inch, 0.6 * inch], repeatRows=1)
        t.setStyle(grid)
        story += [t, Spacer(1, 8)]
    for label, text in context.items():
        if text:
            story += [Paragraph(f"<b>{escape(label)}</b> (recorded, not scored)", body), para(text), Spacer(1, 6)]
    story += [Spacer(1, 10), para(f"Scoring policy: {policy_stamp(result, meta)}", small),
              Paragraph(DISCLAIMER, small), Spacer(1, 22),
              Paragraph("Evaluator signature: ______________________________ &nbsp;&nbsp; Date: ______________", body)]
    doc.build(story)
    return buf.getvalue()
