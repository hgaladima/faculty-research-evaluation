"""Build faculty_research_calculator.xlsx: the scoring engine re-expressed as Excel formulas (no macros).

Run from faculty_research_webapp/:  python -m tools.build_excel_calculator [output.xlsx]

The workbook is generated, never edited by hand. Every setting comes from the provisional policy and every
dropdown label from forms.py, so the calculator, the web app and the engine share one vocabulary. Rule logic
lives in three places that must agree: engine/research_eval.py, forms.build_faculty and the formulas below;
tools/validate_excel.py checks the agreement in real Excel.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import forms
from engine.policy import (EXCEEDS_RULES, GRANT_ROLES, PRES_ROLES, PROVISIONAL, REDUCED_RULES,
                           ROLLING_BENCHMARKS, SUBSTANTIVE_ROLES, fingerprint)
from policy_labels import OPTION_LABELS, POLICY_GROUPS

OUTPUT = Path(__file__).resolve().parents[2] / "faculty_research_calculator.xlsx"
ROWS = {"pubs": 40, "grants": 30, "pres": 30, "editorial": 6, "reviews": 3}
HEADER, FIRST = 4, 5
UNVERIFIED = "Not yet verified"
RANKS = ["Assistant Professor", "Associate Professor", "Professor", "Other"]
APPOINTMENTS = ["Tenure-track", "Tenured", "Non-tenure-track", "Other"]
DOI_PREFIXES = ("https://doi.org/", "http://doi.org/", "https://dx.doi.org/", "http://dx.doi.org/", "doi:")

FONT = "Arial"
NAVY, INK, GREY, LINE = "1D4E89", "14202B", "EEF1F5", "C9D1DA"
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
CALC_FILL = PatternFill("solid", fgColor=GREY)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Border(*(Side(style="thin", color=LINE),) * 4)
UNLOCKED, LOCKED = Protection(locked=False), Protection(locked=True)


def col(i: int) -> str:
    return get_column_letter(i)


def yes(flag: bool) -> str:
    return "Yes" if flag else "No"


class Calculator:
    def __init__(self, policy: dict):
        self.policy = policy
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sheets = {}
        for title in ("Instructions", "Faculty", "Publications", "Grants", "Presentations", "Other",
                      "Dashboard", "Report", "Settings", "Calc", "Lists", "Defaults"):
            self.sheets[title] = self.wb.create_sheet(title)

    # ------------------------------------------------------------------ primitives
    def name(self, name: str, sheet: str, ref: str) -> None:
        absolute = ":".join("$" + "".join(c for c in part if c.isalpha()) + "$" + "".join(c for c in part if c.isdigit())
                            for part in ref.split(":"))
        self.wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!{absolute}")

    def put(self, sheet: str, ref: str, value, *, bold=False, size=10, color=INK, fill=None, wrap=False,
            align=None, fmt=None, border=False, unlocked=False, italic=False):
        cell = self.sheets[sheet][ref]
        cell.value = value
        cell.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
        cell.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
        if fill:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        if border:
            cell.border = THIN
        cell.protection = UNLOCKED if unlocked else LOCKED
        return cell

    def title(self, sheet: str, text: str, sub: str = "") -> None:
        self.put(sheet, "A1", text, bold=True, size=16, color=NAVY)
        if sub:
            self.put(sheet, "A2", sub, color="5A6A79")

    def validate_list(self, sheet: str, ref: str, source: str, prompt: str = "") -> None:
        formula = source if source.startswith('"') else f"={source}"
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorTitle="Choose from the list", error="Pick one of the listed values.")
        if prompt:
            dv.promptTitle, dv.prompt, dv.showInputMessage = "Guidance", prompt[:250], True
        self.sheets[sheet].add_data_validation(dv)
        dv.add(ref)

    def validate_number(self, sheet: str, ref: str, low, high, whole=False, prompt: str = "") -> None:
        dv = DataValidation(type="whole" if whole else "decimal", operator="between", formula1=str(low),
                            formula2=str(high), allow_blank=True, showErrorMessage=True, errorTitle="Check the value",
                            error=f"Enter a {'whole ' if whole else ''}number from {low} to {high}.")
        if prompt:
            dv.promptTitle, dv.prompt, dv.showInputMessage = "Guidance", prompt[:250], True
        self.sheets[sheet].add_data_validation(dv)
        dv.add(ref)

    def protect(self, sheet: str) -> None:
        ws = self.sheets[sheet]
        ws.protection.sheet = True
        ws.protection.formatColumns = False
        ws.protection.formatRows = False

    # ------------------------------------------------------------------ settings and lists
    def build_settings(self) -> None:
        s, p = "Settings", self.policy
        ws = self.sheets[s]
        self.title(s, "Scoring policy settings",
                   "Owned by the annual review committee. The sheet is protected: Review > Unprotect Sheet (no "
                   "password) to change a value. The Dashboard reports how many settings differ from the shipped policy.")
        self.put(s, "A4", "Policy name", bold=True)
        self.put(s, "B4", p["policy_name"], border=True)
        self.name("s_policy_name", s, "B4")
        self.put(s, "A5", "Approval note (who approved, and when)", bold=True)
        self.put(s, "B5", "Provisional defaults; not approved", border=True)
        self.name("s_policy_note", s, "B5")
        r = 7
        for group, _, fields in POLICY_GROUPS:
            self.put(s, f"A{r}", group, bold=True, color="FFFFFF", fill=HEAD_FILL)
            self.put(s, f"B{r}", "Value", bold=True, color="FFFFFF", fill=HEAD_FILL)
            self.put(s, f"C{r}", "Meaning of the value", bold=True, color="FFFFFF", fill=HEAD_FILL)
            r += 1
            for path, label, kind in fields:
                if isinstance(kind, tuple):
                    continue
                node = p
                for part in path.split("."):
                    node = node[part]
                key = "s_" + path.replace(".", "_")
                self.put(s, f"A{r}", label, wrap=True)
                if path == "rounding":
                    self.put(s, f"B{r}", "half_up", border=True)
                    self.put(s, f"C{r}", "Fixed in Excel: scores round half up. The Phase 1 'round half to even' "
                                         "rule is available only in the web app.", wrap=True, color="5A6A79")
                else:
                    value = yes(node) if isinstance(node, bool) else node
                    self.put(s, f"B{r}", value, border=True)
                    if kind == "bool":
                        self.validate_list(s, f"B{r}", "L_yesno")
                    elif isinstance(kind, list):
                        self.validate_list(s, f"B{r}", '"' + ",".join(str(k) for k in kind) + '"')
                        if isinstance(node, str) and node in OPTION_LABELS:
                            options = "; ".join(f"{k} = {OPTION_LABELS[k]}" for k in kind if k in OPTION_LABELS)
                            self.put(s, f"C{r}", options, wrap=True, color="5A6A79")
                self.name(key, s, f"B{r}")
                r += 1
            r += 1

        def table(title: str, headers: list[str], rows: list[list], names: list[str | None]) -> None:
            nonlocal r
            self.put(s, f"A{r}", title, bold=True, color="FFFFFF", fill=HEAD_FILL)
            for i, h in enumerate(headers):
                self.put(s, f"{col(i + 1)}{r + 1}", h, bold=True, fill=CALC_FILL, border=True, wrap=True)
            for j, row in enumerate(rows):
                for i, value in enumerate(row):
                    self.put(s, f"{col(i + 1)}{r + 2 + j}", value, border=True)
            for i, n in enumerate(names):
                if n:
                    self.name(n, s, f"{col(i + 1)}{r + 2}:{col(i + 1)}{r + 1 + len(rows)}")
            r += len(rows) + 3

        tiers = [(label, tier) for label, tier in forms.JOURNAL_TIERS.items() if tier is not None]
        table("Journal tiers", ["Tier", "Number", "Multiplier"],
              [[label, tier, p["journal_tier_mult"][tier]] for label, tier in tiers], ["TR_label", "TR_num", "TR_mult"])
        roles = p["author_role_mult"]
        table("Author role multipliers", ["Role", "Multiplier"],
              [["lead", roles["lead"]], ["second", roles.get("second", roles["coauthor"])],
               ["coauthor", roles["coauthor"]]], ["RM_code", "RM_mult"])
        table("Authorship choices", ["Choice on the entry sheet", "Scored as"],
              [[label, code] for label, code in forms.AUTHORSHIP.items() if code], ["AU_label", "AU_code"])
        table("Publication types", ["Choice on the entry sheet", "Code"],
              [[label, code] for label, code in forms.PUB_TYPES.items()], ["PT_label", "PT_code"])
        other = dict(p["other_output_units"])
        for key in ("non_peer_reviewed_book", "non_peer_reviewed_chapter", "published_abstract"):
            other.setdefault(key, 0.0)
        table("Other scholarly outputs (units each; 0 = not credited)", ["Code", "Units"],
              [[k, v] for k, v in other.items()], ["OO_key", "OO_units"])
        table("Editorial and peer-review service (units each)", ["Role", "Units"],
              [[label, p["editorial_units"][code]] for label, code in forms.EDITORIAL_ROLES.items()],
              ["ED_label", "ED_units"])
        table("Grant roles", ["Role", "Code", "Multiplier", "Substantive role", "A new award meets the grant minimum"],
              [[label, code, p["grant_role_mult"][code], yes(code in SUBSTANTIVE_ROLES),
                yes(code in p["award_gate_roles"])] for label, code in forms.GRANT_ROLES.items() if code],
              ["GR_label", "GR_code", "GR_mult", "GR_subst", "GR_gate"])
        table("Grant size bands (units)", ["Band", "Proposal submitted", "Award received"],
              [[band, p["submission_base"][band], p["award_base"][band]] for band in ("internal", "minor", "major", "large")],
              ["SZ_band", "SZ_sub", "SZ_award"])
        table("Grant review panels (units each)", ["Panel type", "Units"],
              [[label, p["grant_review_units"][code]] for label, code in forms.REVIEW_TYPES.items()],
              ["RV_label", "RV_units"])
        table("Presentation types", ["Type", "Base units", "Local dissemination"],
              [[label, p["pres_type_base"][code], yes(code in forms.LOCAL_PRES)]
               for label, code in forms.PRES_TYPES.items()], ["PR_label", "PR_base", "PR_local"])
        table("Venue tiers", ["Tier", "Code", "Multiplier"],
              [[label, code, p["venue_tier_mult"][code]] for label, code in forms.VENUE_TIERS.items() if code],
              ["VN_label", "VN_code", "VN_mult"])
        table("Presentation roles", ["Role", "Code", "Multiplier", "Meets the presentation minimum"],
              [[label, code, p["pres_role_mult"][code], yes(code in p["pres_gate_roles"])]
               for label, code in forms.PRES_ROLES.items() if code], ["RO_label", "RO_code", "RO_mult", "RO_gate"])
        self.settings_last_row = r
        for letter, width in zip("ABCDE", (62, 30, 60, 18, 22)):
            ws.column_dimensions[letter].width = width
        ws.freeze_panes = "A4"

        mirror = self.sheets["Defaults"]
        for row in ws.iter_rows(min_row=1, max_row=r, max_col=8):
            for cell in row:
                if cell.value is not None:
                    mirror[cell.coordinate] = cell.value
        self.put("Defaults", "J1", "Shipped values of the Settings sheet. Do not edit.", italic=True)
        self.put("Defaults", "J2", fingerprint(p))
        self.name("d_fingerprint", "Defaults", "J2")
        mirror.sheet_state = "hidden"

    def build_lists(self) -> None:
        lists = {
            "L_yesno": ["Yes", "No"], "L_rank": RANKS, "L_appointment": APPOINTMENTS,
            "L_pubtype": list(forms.PUB_TYPES), "L_tier": list(forms.JOURNAL_TIERS),
            "L_authorship": list(forms.AUTHORSHIP), "L_quartile": forms.QUARTILES,
            "L_source": forms.GRANT_SOURCES, "L_grantrole": list(forms.GRANT_ROLES), "L_status": forms.GRANT_STATUSES,
            "L_prestype": list(forms.PRES_TYPES), "L_venue": list(forms.VENUE_TIERS),
            "L_presrole": list(forms.PRES_ROLES), "L_scope": forms.PRES_SCOPES,
            "L_editorial": list(forms.EDITORIAL_ROLES), "L_review": list(forms.REVIEW_TYPES),
        }
        for i, (name, values) in enumerate(lists.items(), 1):
            self.put("Lists", f"{col(i)}1", name, bold=True)
            for j, value in enumerate(values, 2):
                self.put("Lists", f"{col(i)}{j}", value)
            self.name(name, "Lists", f"{col(i)}2:{col(i)}{len(values) + 1}")
        self.sheets["Lists"].sheet_state = "hidden"

    # ------------------------------------------------------------------ entry sheets
    def entry_table(self, sheet: str, title: str, sub: str, columns: list[tuple], n: int) -> None:
        """columns: (header, width, kind, source, prompt). kind: text | list | yesno | whole | money | number."""
        ws = self.sheets[sheet]
        self.title(sheet, title, sub)
        self.put(sheet, "A3", "Yellow cells are for entry. Grey columns to the right are calculated; a message in "
                              "'Check this row' means the row is not counted until it is completed.", color="5A6A79")
        last = FIRST + n - 1
        for i, (header, width, kind, source, prompt) in enumerate(columns, 1):
            letter = col(i)
            self.put(sheet, f"{letter}{HEADER}", header, bold=True, color="FFFFFF", fill=HEAD_FILL, wrap=True, border=True)
            ws.column_dimensions[letter].width = width
            for r in range(FIRST, last + 1):
                fmt = "$#,##0" if kind == "money" else "0" if kind == "whole" else None
                self.put(sheet, f"{letter}{r}", None, fill=INPUT_FILL, border=True, unlocked=True, fmt=fmt)
            ref = f"{letter}{FIRST}:{letter}{last}"
            if kind == "list":
                self.validate_list(sheet, ref, source, prompt)
            elif kind == "yesno":
                self.validate_list(sheet, ref, "L_yesno", prompt)
            elif kind == "whole":
                self.validate_number(sheet, ref, 0, 999, whole=True, prompt=prompt)
            elif kind == "money":
                self.validate_number(sheet, ref, 0, 10 ** 10, prompt=prompt)
            elif kind == "number":
                self.validate_number(sheet, ref, 0, 1000, prompt=prompt)
        ws.freeze_panes = f"B{FIRST}"
        ws.row_dimensions[HEADER].height = 44

    def calc_columns(self, sheet: str, start: int, n: int, specs: list[tuple]) -> dict:
        """specs: (key, header, formula(r, c) -> str). Returns key -> column letter; names each column <prefix>_<key>."""
        ws = self.sheets[sheet]
        letters = {key: col(start + i) for i, (key, _, _) in enumerate(specs)}
        last = FIRST + n - 1
        self.put(sheet, f"{col(start)}3", "Calculated (do not edit)", bold=True, color="5A6A79")
        for key, header, formula in specs:
            letter = letters[key]
            self.put(sheet, f"{letter}{HEADER}", header, bold=True, fill=CALC_FILL, wrap=True, border=True)
            wide = key in ("problem", "note")
            ws.column_dimensions[letter].width = 46 if wide else 11
            for r in range(FIRST, last + 1):
                self.put(sheet, f"{letter}{r}", "=" + formula(r, letters), fill=CALC_FILL, border=True,
                         color="B0362D" if key == "problem" else INK)
            self.name(f"{sheet[:3].upper()}_{key}", sheet, f"{letter}{FIRST}:{letter}{last}")
        return letters

    def build_publications(self) -> None:
        sh, n = "Publications", ROWS["pubs"]
        tier_help = ("Tier 1: top quartile (Q1) of its subject category, or the flagship journal of a major society. "
                     "Tier 2: second quartile. Tier 3: other indexed peer-reviewed journal.")
        self.entry_table(sh, "Publications and scholarly outputs",
                         "One row per item first published in the evaluation year. An item counts once.", [
            ("Title", 34, "text", None, ""), ("Journal or publisher", 24, "text", None, ""),
            ("Type", 22, "list", "L_pubtype", ""), ("Peer reviewed", 10, "yesno", None, "Blank counts as Yes."),
            ("Journal tier", 22, "list", "L_tier", tier_help),
            ("Authorship", 22, "list", "L_authorship", "Sole, first, senior (last) and corresponding authors receive lead credit."),
            ("Student or mentee first author", 12, "yesno", None, "Yes when a student or mentee is first author and you are senior or corresponding."),
            ("Documented impacts", 11, "whole", None, "Documented uses in policy, guidelines or practice."),
            ("DOI", 22, "text", None, "Used to stop the same item being counted twice."),
            ("Quartile (info)", 10, "list", "L_quartile", "Recorded; not scored."),
            ("Impact factor (info)", 10, "number", None, "Recorded; not scored."),
            ("Credited in a prior year", 11, "yesno", None, "Yes removes this item from this year's credit."),
        ], n)
        first, last = FIRST, FIRST + n - 1
        strip = "{raw}"
        for prefix in reversed(DOI_PREFIXES):
            strip = f'IF(LEFT({{raw}},{len(prefix)})="{prefix}",MID({{raw}},{len(prefix) + 1},500),{strip})'

        def kept(r, c):
            best = (f'SUMPRODUCT(({c["doi"]}${first}:{c["doi"]}${last}={c["doi"]}{r})*'
                    f'({c["expected"]}${first}:{c["expected"]}${last}>{c["expected"]}{r}))=0')
            if r == first:
                return f'IF({c["doi"]}{r}="",1,IF({best},1,0))'
            earlier = (f'SUMPRODUCT(({c["doi"]}${first}:{c["doi"]}{r - 1}={c["doi"]}{r})*'
                       f'({c["expected"]}${first}:{c["expected"]}{r - 1}={c["expected"]}{r}))=0')
            return f'IF({c["doi"]}{r}="",1,IF(AND({best},{earlier}),1,0))'

        book = lambda r, c: f's_book_units*IF({c["role"]}{r}="lead",1,s_book_coauthor_mult)'
        specs = [
            ("used", "Row used", lambda r, c: f'IF(C{r}<>"",1,0)'),
            ("code", "Type code", lambda r, c: f'IF({c["used"]}{r}=0,"",INDEX(PT_code,MATCH(C{r},PT_label,0)))'),
            ("peer", "Peer reviewed", lambda r, c: f'IF(D{r}="No",0,1)'),
            ("problem", "Check this row", lambda r, c:
                f'IF(AND(C{r}="",OR(A{r}<>"",B{r}<>"")),"Choose a publication type",IF({c["used"]}{r}=0,"",'
                f'IF(AND({c["code"]}{r}="article",{c["peer"]}{r}=1,E{r}=""),"Choose a journal tier (or \'Not yet verified\')",'
                f'IF(AND(OR({c["code"]}{r}="article",{c["code"]}{r}="book"),F{r}=""),'
                f'"Choose the authorship (or \'Not yet verified\')",""))))'),
            ("ok", "Counted row", lambda r, c: f'IF(AND({c["used"]}{r}=1,{c["problem"]}{r}=""),1,0)'),
            ("prior", "Prior year", lambda r, c:
                f'IF(AND({c["ok"]}{r}=1,s_honour_prior_year_publications="Yes",L{r}="Yes"),1,0)'),
            ("tier", "Tier", lambda r, c:
                f'IF(OR(E{r}="",E{r}="{UNVERIFIED}"),3,INDEX(TR_num,MATCH(E{r},TR_label,0)))'),
            ("tiermult", "Tier multiplier", lambda r, c: f'INDEX(TR_mult,MATCH({c["tier"]}{r},TR_num,0))'),
            ("role", "Role", lambda r, c:
                f'IF(OR(F{r}="",F{r}="{UNVERIFIED}"),IF({c["code"]}{r}="book",s_missing_defaults_book_role,"coauthor"),'
                f'INDEX(AU_code,MATCH(F{r},AU_label,0)))'),
            ("rolemult", "Role multiplier", lambda r, c: f'INDEX(RM_mult,MATCH({c["role"]}{r},RM_code,0))'),
            ("kind", "Scored as", lambda r, c:
                f'IF({c["ok"]}{r}=0,"",IF(AND({c["code"]}{r}="article",{c["peer"]}{r}=1),IF({c["tier"]}{r}=0,"tier0","article"),'
                f'IF(AND({c["code"]}{r}="book",{c["peer"]}{r}=1),"book",IF(AND({c["code"]}{r}="chapter",{c["peer"]}{r}=1),"chapter","other"))))'),
            ("otherunits", "Other-output units", lambda r, c:
                f'IF({c["kind"]}{r}<>"other",0,IFERROR(INDEX(OO_units,MATCH(IF(OR({c["code"]}{r}="article",{c["code"]}{r}="book",'
                f'{c["code"]}{r}="chapter"),"non_peer_reviewed_"&{c["code"]}{r},{c["code"]}{r}),OO_key,0)),0))'),
            ("expected", "Stand-alone units", lambda r, c:
                f'IF({c["ok"]}{r}=0,0,ROUND(IF({c["prior"]}{r}=1,-1,IF({c["kind"]}{r}="article",{c["tiermult"]}{r}*{c["rolemult"]}{r},'
                f'IF({c["kind"]}{r}="tier0",0,IF({c["kind"]}{r}="book",{book(r, c)},IF({c["kind"]}{r}="chapter",s_chapter_units,'
                f'{c["otherunits"]}{r}))))),6))'),
            ("doiraw", "DOI as typed", lambda r, c: f'IF({c["ok"]}{r}=0,"",LOWER(TRIM(I{r})))'),
            ("doi", "DOI compared", lambda r, c:
                f'IF({c["doiraw"]}{r}="","",IF(s_normalise_ids="Yes",{strip.format(raw=c["doiraw"] + str(r))},{c["doiraw"]}{r}))'),
            ("kept", "Copy kept", kept),
            ("dup", "Duplicate", lambda r, c:
                f'IF(AND({c["ok"]}{r}=1,{c["doi"]}{r}<>"",{c["kept"]}{r}=0,s_dedupe_doi="Yes"),1,0)'),
            ("live", "Credited", lambda r, c: f'IF(AND({c["ok"]}{r}=1,{c["prior"]}{r}=0,{c["dup"]}{r}=0),1,0)'),
            ("main", "Article and book units", lambda r, c:
                f'IF({c["live"]}{r}=0,0,IF({c["kind"]}{r}="article",{c["tiermult"]}{r}*{c["rolemult"]}{r}'
                f'+IF(AND(G{r}="Yes",{c["role"]}{r}="lead"),s_student_mentee_increment,0),IF({c["kind"]}{r}="book",{book(r, c)},0)))'),
            ("chapter", "Chapter units", lambda r, c: f'IF(AND({c["live"]}{r}=1,{c["kind"]}{r}="chapter"),s_chapter_units,0)'),
            ("other", "Other units", lambda r, c: f'IF(AND({c["live"]}{r}=1,{c["kind"]}{r}="other"),{c["otherunits"]}{r},0)'),
            ("ischap", "Chapter", lambda r, c: f'IF(AND({c["live"]}{r}=1,{c["kind"]}{r}="chapter"),1,0)'),
            ("count", "Counts toward minimum", lambda r, c:
                f'IF({c["live"]}{r}=0,0,IF({c["kind"]}{r}="article",1,IF({c["kind"]}{r}="book",2,'
                f'IF(AND({c["kind"]}{r}="chapter",SUM({c["ischap"]}${HEADER}:{c["ischap"]}{r - 1})=0),1,0))))'),
            ("impact", "Impact units", lambda r, c:
                f'IF(AND({c["live"]}{r}=1,{c["kind"]}{r}<>"tier0",NOT(AND({c["code"]}{r}="article",{c["peer"]}{r}=0))),'
                f's_impact_units*N(H{r}),0)'),
            ("note", "Note", lambda r, c:
                f'IF({c["ok"]}{r}=0,"",IF({c["prior"]}{r}=1,"Credited in a prior year: excluded",IF({c["dup"]}{r}=1,'
                f'"Duplicate DOI: counted once (the entry with the most credit is kept)",IF({c["kind"]}{r}="tier0",'
                f'"Questionable venue: not credited",IF(AND({c["kind"]}{r}="article",OR(E{r}="",E{r}="{UNVERIFIED}")),'
                f'"Journal tier not verified: treated as Tier 3",IF(AND(OR({c["kind"]}{r}="article",{c["kind"]}{r}="book"),'
                f'F{r}="{UNVERIFIED}"),"Author role not verified: treated as "&{c["role"]}{r},'
                f'IF(AND({c["kind"]}{r}="other",{c["otherunits"]}{r}=0),"This type earns no credit under the current settings","")))))))'),
        ]
        self.calc_columns(sh, 14, n, specs)
        self.protect(sh)

    def build_grants(self) -> None:
        sh, n = "Grants", ROWS["grants"]
        sub, awd, act = forms.SUBMITTED, forms.AWARDED, forms.ACTIVE
        self.entry_table(sh, "Grants and contracts",
                         "One row per proposal, award or continuing project. For a proposal submitted and funded in the "
                         "same year, enter one row as 'Awarded this year' and answer Yes to 'Also submitted this year'.", [
            ("Title", 34, "text", None, ""), ("Funder", 22, "text", None, ""),
            ("Source", 11, "list", "L_source", "Blank counts as External."),
            ("Total amount ($)", 15, "money", None, "Total for the full project period. Not needed for internal funding or active projects."),
            ("Role", 16, "list", "L_grantrole", ""),
            ("Competitive", 11, "yesno", None, "Blank counts as Yes. No for sole-source contracts."),
            ("Status", 30, "list", "L_status", ""),
            ("Also submitted this year", 11, "yesno", None, "Only for 'Awarded this year'."),
            ("Scored or discussed", 11, "yesno", None, "The proposal received a score or was discussed by the panel."),
            ("Credited in a prior year", 11, "yesno", None, "Yes removes this row from this year's credit."),
        ], n)
        by_code = lambda column, r, c: f'INDEX({column},MATCH({c["rolecode"]}{r},GR_code,0))'
        size_lookup = lambda column, r, c: f'INDEX({column},MATCH({c["size"]}{r},SZ_band,0))'
        specs = [
            ("used", "Row used", lambda r, c: f'IF(G{r}<>"",1,0)'),
            ("problem", "Check this row", lambda r, c:
                f'IF(AND(G{r}="",OR(A{r}<>"",B{r}<>"")),"Choose a status and a role",IF({c["used"]}{r}=0,"",'
                f'IF(E{r}="","Choose a role (or \'Not yet verified\')",IF(AND(G{r}="{act}",H{r}="Yes"),'
                f'"A project awarded in a prior year cannot also be \'submitted this year\'. Add a separate row for the new submission",'
                f'IF(AND(G{r}<>"{act}",C{r}<>"Internal",NOT(ISNUMBER(D{r}))),"Enter the total amount","")))))'),
            ("ok", "Counted row", lambda r, c: f'IF(AND({c["used"]}{r}=1,{c["problem"]}{r}=""),1,0)'),
            ("rolecode", "Role code", lambda r, c:
                f'IF({c["ok"]}{r}=0,"",IF(E{r}="{UNVERIFIED}","collaborator",INDEX(GR_code,MATCH(E{r},GR_label,0))))'),
            ("rolemult", "Role multiplier", lambda r, c: f'IF({c["ok"]}{r}=0,0,{by_code("GR_mult", r, c)})'),
            ("subst", "Substantive role", lambda r, c: f'IF({c["ok"]}{r}=0,0,IF({by_code("GR_subst", r, c)}="Yes",1,0))'),
            ("gaterole", "Award role meets minimum", lambda r, c:
                f'IF({c["ok"]}{r}=0,0,IF({by_code("GR_gate", r, c)}="Yes",1,0))'),
            ("size", "Size band", lambda r, c:
                f'IF(OR({c["ok"]}{r}=0,G{r}="{act}"),"",IF(C{r}="Internal","internal",IF(D{r}<s_amount_bands_minor_max,"minor",'
                f'IF(D{r}<s_amount_bands_large_min,"major","large"))))'),
            ("comp", "Competitive", lambda r, c: f'IF(F{r}="No",0,1)'),
            ("prior", "Prior year", lambda r, c: f'IF(J{r}="Yes",1,0)'),
            ("issub", "Submission", lambda r, c:
                f'IF(AND({c["ok"]}{r}=1,G{r}<>"{act}",OR(G{r}="{sub}",H{r}="Yes"),{c["prior"]}{r}=0),1,0)'),
            ("isaward", "Award", lambda r, c: f'IF(AND({c["ok"]}{r}=1,G{r}="{awd}",{c["prior"]}{r}=0),1,0)'),
            ("meets", "Submission meets minimum", lambda r, c:
                f'IF(AND({c["issub"]}{r}=1,{c["subst"]}{r}=1,{c["comp"]}{r}=1,OR({c["size"]}{r}<>"internal",'
                f's_grant_gate_requires_external="No")),1,0)'),
            ("deduped", "Credited as award only", lambda r, c:
                f'IF(AND({c["issub"]}{r}=1,{c["isaward"]}{r}=1,s_dedupe_submission_award="Yes"),1,0)'),
            ("subunits", "Submission units", lambda r, c:
                f'IF(AND({c["issub"]}{r}=1,{c["deduped"]}{r}=0),{size_lookup("SZ_sub", r, c)}*{c["rolemult"]}{r}'
                f'*IF({c["comp"]}{r}=1,1,s_noncompetitive_mult)+IF(I{r}="Yes",s_scored_bonus*{c["rolemult"]}{r},0),0)'),
            ("subpartial", "Submission below minimum", lambda r, c:
                f'IF(AND({c["issub"]}{r}=1,{c["deduped"]}{r}=0,{c["meets"]}{r}=0),1,0)'),
            ("awardunits", "Award units", lambda r, c:
                f'IF({c["isaward"]}{r}=1,{size_lookup("SZ_award", r, c)}*{c["rolemult"]}{r}'
                f'*IF({c["comp"]}{r}=1,1,s_noncompetitive_award_mult),0)'),
            ("awardgate", "Award meets minimum", lambda r, c: f'IF(AND({c["isaward"]}{r}=1,{c["gaterole"]}{r}=1),1,0)'),
            ("awardpartial", "Award below minimum", lambda r, c: f'IF(AND({c["isaward"]}{r}=1,{c["gaterole"]}{r}=0),1,0)'),
            ("distinct", "Large competitive award as PI/MPI", lambda r, c:
                f'IF(AND({c["isaward"]}{r}=1,{c["size"]}{r}="large",OR({c["rolecode"]}{r}="PI",{c["rolecode"]}{r}="MPI"),'
                f'{c["comp"]}{r}=1),1,0)'),
            ("active", "Active project credited", lambda r, c: f'IF(AND({c["ok"]}{r}=1,G{r}="{act}",{c["subst"]}{r}=1),1,0)'),
            ("note", "Note", lambda r, c:
                f'IF({c["ok"]}{r}=0,"",IF({c["prior"]}{r}=1,"Credited in a prior year: no credit this year",'
                f'IF(E{r}="{UNVERIFIED}","Role not verified: treated as collaborator",IF({c["deduped"]}{r}=1,'
                f'"Funded in the same year: credited as an award only",IF(AND(G{r}="{act}",{c["subst"]}{r}=0),'
                f'"Not a substantive role: no implementation credit","")))))'),
        ]
        self.calc_columns(sh, 12, n, specs)
        self.protect(sh)

    def build_presentations(self) -> None:
        sh, n = "Presentations", ROWS["pres"]
        venue_help = ("Tier A: major national or international society meeting with competitive abstract review. "
                      "Tier B: regional or state society meeting. Tier C: local or limited review.")
        self.entry_table(sh, "Scientific presentations",
                         "One row per presentation. Venue tier is not used for institutional, outreach or seminar talks.", [
            ("Title", 34, "text", None, ""), ("Meeting", 26, "text", None, ""),
            ("Location (info)", 16, "text", None, "Recorded; not scored."),
            ("Scope (info)", 14, "list", "L_scope", "Recorded; not scored."),
            ("Type", 30, "list", "L_prestype", ""), ("Venue tier", 24, "list", "L_venue", venue_help),
            ("Role", 26, "list", "L_presrole", ""),
            ("Peer-reviewed abstract (info)", 12, "yesno", None, "Recorded; not scored."),
        ], n)
        by_code = lambda column, r, c: f'INDEX({column},MATCH({c["rolecode"]}{r},RO_code,0))'
        specs = [
            ("used", "Row used", lambda r, c: f'IF(E{r}<>"",1,0)'),
            ("local", "Local", lambda r, c:
                f'IF({c["used"]}{r}=0,0,IF(INDEX(PR_local,MATCH(E{r},PR_label,0))="Yes",1,0))'),
            ("problem", "Check this row", lambda r, c:
                f'IF(AND(E{r}="",OR(A{r}<>"",B{r}<>"")),"Choose a presentation type and a role",IF({c["used"]}{r}=0,"",'
                f'IF(G{r}="","Choose a role (or \'Not yet verified\')",IF(AND({c["local"]}{r}=0,F{r}=""),'
                f'"Choose a venue tier (or \'Not yet verified\')",""))))'),
            ("ok", "Counted row", lambda r, c: f'IF(AND({c["used"]}{r}=1,{c["problem"]}{r}=""),1,0)'),
            ("base", "Base units", lambda r, c: f'IF({c["ok"]}{r}=0,0,INDEX(PR_base,MATCH(E{r},PR_label,0)))'),
            ("rolecode", "Role code", lambda r, c:
                f'IF({c["ok"]}{r}=0,"",IF(G{r}="{UNVERIFIED}",s_missing_defaults_pres_role,INDEX(RO_code,MATCH(G{r},RO_label,0))))'),
            ("rolemult", "Role multiplier", lambda r, c: f'IF({c["ok"]}{r}=0,0,{by_code("RO_mult", r, c)})'),
            ("rolegate", "Role meets minimum", lambda r, c: f'IF({c["ok"]}{r}=0,0,IF({by_code("RO_gate", r, c)}="Yes",1,0))'),
            ("venue", "Venue code", lambda r, c:
                f'IF(OR({c["ok"]}{r}=0,{c["local"]}{r}=1),"",IF(F{r}="{UNVERIFIED}",s_missing_defaults_venue,'
                f'INDEX(VN_code,MATCH(F{r},VN_label,0))))'),
            ("venuemult", "Venue multiplier", lambda r, c:
                f'IF({c["venue"]}{r}="",0,INDEX(VN_mult,MATCH({c["venue"]}{r},VN_code,0)))'),
            ("localunits", "Local units", lambda r, c:
                f'IF(AND({c["ok"]}{r}=1,{c["local"]}{r}=1),{c["base"]}{r}*{c["rolemult"]}{r},0)'),
            ("units", "Conference units", lambda r, c:
                f'IF(AND({c["ok"]}{r}=1,{c["local"]}{r}=0,{c["venue"]}{r}<>"0"),{c["base"]}{r}*{c["venuemult"]}{r}*{c["rolemult"]}{r},0)'),
            ("gate", "Meets minimum", lambda r, c:
                f'IF(AND({c["ok"]}{r}=1,{c["local"]}{r}=0,{c["venue"]}{r}<>"0",{c["rolegate"]}{r}=1),1,0)'),
            ("partial", "Partial activity", lambda r, c:
                f'IF({c["ok"]}{r}=0,0,IF({c["local"]}{r}=1,1,IF(AND({c["venue"]}{r}<>"0",{c["rolegate"]}{r}=0),1,0)))'),
            ("note", "Note", lambda r, c:
                f'IF({c["ok"]}{r}=0,"",IF(AND({c["local"]}{r}=0,{c["venue"]}{r}="0"),"Questionable venue: not credited",'
                f'IF(G{r}="{UNVERIFIED}","Role not verified: treated as "&{c["rolecode"]}{r},IF(AND({c["local"]}{r}=0,'
                f'F{r}="{UNVERIFIED}"),"Venue tier not verified: treated as Tier "&{c["venue"]}{r},""))))'),
        ]
        self.calc_columns(sh, 10, n, specs)
        self.protect(sh)

    def build_faculty_and_other(self) -> None:
        sh = "Faculty"
        self.title(sh, "Faculty information", "Yellow cells are for entry.")
        fields = [("Faculty member name", "f_name", "text"), ("Faculty ID", "f_id", "text"),
                  ("Academic rank", "f_rank", "L_rank"), ("Department", "f_department", "text"),
                  ("Appointment type", "f_appointment", "L_appointment"), ("Evaluation year", "f_year", "text")]
        for i, (label, name, kind) in enumerate(fields):
            self.put(sh, f"A{4 + i}", label, bold=True)
            self.put(sh, f"B{4 + i}", None, fill=INPUT_FILL, border=True, unlocked=True, fmt="@" if name == "f_year" else None)
            self.name(name, sh, f"B{4 + i}")
            if kind.startswith("L_"):
                self.validate_list(sh, f"B{4 + i}", kind)
        self.put(sh, "A11", "Research assignment and adjustments", bold=True, color=NAVY, size=12)
        rows = [("Assigned research effort (%)", "f_effort", lambda ref: self.validate_number(
                    sh, ref, 0, 100, prompt="Leave blank to use the policy's standard research assignment.")),
                ("Months active in the evaluation period", "f_months", lambda ref: self.validate_number(
                    sh, ref, 0, 12, prompt="Blank counts as 12. Subtract approved non-research leave; count sabbatical as active.")),
                ("First-year faculty member", "f_first_year", lambda ref: self.validate_list(sh, ref, "L_yesno"))]
        for i, (label, name, validate) in enumerate(rows):
            self.put(sh, f"A{12 + i}", label, bold=True)
            self.put(sh, f"B{12 + i}", None, fill=INPUT_FILL, border=True, unlocked=True)
            self.name(name, sh, f"B{12 + i}")
            validate(f"B{12 + i}")
        self.put(sh, "C12", '=IF(f_effort="","Using the standard assignment of "&s_standard_research_effort&"%","")', color="5A6A79")
        self.sheets[sh].column_dimensions["A"].width = 42
        self.sheets[sh].column_dimensions["B"].width = 36
        self.sheets[sh].column_dimensions["C"].width = 50
        self.protect(sh)

        sh = "Other"
        self.title(sh, "Other scholarly activities", "Credited inside the three domains, with caps. Yellow cells are for entry.")
        ne, nr = ROWS["editorial"], ROWS["reviews"]
        for letter, header in zip("ABC", ("Editorial and peer-review role", "Count", "Units")):
            self.put(sh, f"{letter}{HEADER}", header, bold=True, color="FFFFFF", fill=HEAD_FILL, border=True, wrap=True)
        for letter, header in zip("EFG", ("Grant review panel", "Count", "Units")):
            self.put(sh, f"{letter}{HEADER}", header, bold=True, color="FFFFFF", fill=HEAD_FILL, border=True, wrap=True)
        for r in range(FIRST, FIRST + ne):
            self.put(sh, f"A{r}", None, fill=INPUT_FILL, border=True, unlocked=True)
            self.put(sh, f"B{r}", None, fill=INPUT_FILL, border=True, unlocked=True, fmt="0")
            self.put(sh, f"C{r}", f'=IF(A{r}="",0,INDEX(ED_units,MATCH(A{r},ED_label,0))*MAX(N(B{r}),1))', fill=CALC_FILL, border=True)
        for r in range(FIRST, FIRST + nr):
            self.put(sh, f"E{r}", None, fill=INPUT_FILL, border=True, unlocked=True)
            self.put(sh, f"F{r}", None, fill=INPUT_FILL, border=True, unlocked=True, fmt="0")
            self.put(sh, f"G{r}", f'=IF(E{r}="",0,INDEX(RV_units,MATCH(E{r},RV_label,0))*MAX(N(F{r}),1))', fill=CALC_FILL, border=True)
        self.validate_list(sh, f"A{FIRST}:A{FIRST + ne - 1}", "L_editorial")
        self.validate_list(sh, f"E{FIRST}:E{FIRST + nr - 1}", "L_review")
        self.validate_number(sh, f"B{FIRST}:B{FIRST + ne - 1}", 1, 999, whole=True, prompt="Blank counts as 1.")
        self.validate_number(sh, f"F{FIRST}:F{FIRST + nr - 1}", 1, 99, whole=True, prompt="Blank counts as 1.")
        self.name("OT_editorial", sh, f"C{FIRST}:C{FIRST + ne - 1}")
        self.name("OT_review", sh, f"G{FIRST}:G{FIRST + nr - 1}")
        r = FIRST + ne + 2
        self.put(sh, f"A{r}", "Counts for the year", bold=True, color=NAVY, size=12)
        counts = [("Conference sessions chaired or organised", "o_chaired"),
                  ("Grant proposals in development (documented)", "o_pid"),
                  ("Unfunded research projects under way", "o_unfunded"),
                  ("Manuscripts under review", "o_manuscripts"), ("Conference abstracts submitted", "o_abstracts")]
        for i, (label, name) in enumerate(counts, 1):
            self.put(sh, f"A{r + i}", label)
            self.put(sh, f"B{r + i}", None, fill=INPUT_FILL, border=True, unlocked=True, fmt="0")
            self.name(name, sh, f"B{r + i}")
            self.validate_number(sh, f"B{r + i}", 0, 999, whole=True)
        r += len(counts) + 2
        self.put(sh, f"A{r}", "Three-year rolling window (used only when a minimum is unmet this year)", bold=True, color=NAVY, size=12)
        self.put(sh, f"A{r + 1}", "Leave a cell blank when the figure is not provided. A 0 is treated as provided.", color="5A6A79")
        rolling = [("Publications counted in the two prior years", "o_pubs_prior2", 999),
                   ("Prior two years with the grant minimum met (0-2)", "o_grant_prior2", 2),
                   ("Qualifying presentations in the two prior years", "o_pres_prior2", 999)]
        for i, (label, name, high) in enumerate(rolling, 2):
            self.put(sh, f"A{r + i}", label)
            self.put(sh, f"B{r + i}", None, fill=INPUT_FILL, border=True, unlocked=True, fmt="0")
            self.name(name, sh, f"B{r + i}")
            self.validate_number(sh, f"B{r + i}", 0, high, whole=True)
        r += len(rolling) + 3
        self.put(sh, f"A{r}", "Recorded for the committee and the narrative (not scored)", bold=True, color=NAVY, size=12)
        context = [("Research awards and recognition", "x_awards"),
                   ("Research mentoring and dissertation supervision", "x_mentoring"),
                   ("Community-engaged, translational and interdisciplinary work", "x_community"),
                   ("Evaluator notes", "x_notes")]
        for i, (label, name) in enumerate(context):
            row = r + 1 + i
            self.put(sh, f"A{row}", label)
            self.sheets[sh].merge_cells(f"B{row}:G{row}")
            self.put(sh, f"B{row}", None, fill=INPUT_FILL, border=True, unlocked=True, wrap=True)
            self.sheets[sh].row_dimensions[row].height = 48
            self.name(name, sh, f"B{row}")
        for letter, width in zip("ABCDEFG", (50, 12, 10, 3, 26, 10, 10)):
            self.sheets[sh].column_dimensions[letter].width = width
        self.protect(sh)

    # ------------------------------------------------------------------ calculations
    def build_calc(self) -> None:
        sh = "Calc"
        self.title(sh, "Calculations", "Each line mirrors a step of the scoring engine. Do not edit.")
        for letter, header, width in zip("ABC", ("Step", "Value", "Name"), (64, 22, 22)):
            self.put(sh, f"{letter}3", header, bold=True, color="FFFFFF", fill=HEAD_FILL)
            self.sheets[sh].column_dimensions[letter].width = width
        self.calc_row = 4

        def step(name: str, label: str, formula: str) -> None:
            r = self.calc_row
            self.put(sh, f"A{r}", label, wrap=True)
            self.put(sh, f"B{r}", "=" + formula, fill=CALC_FILL, border=True)
            self.put(sh, f"C{r}", name, color="5A6A79")
            self.name(name, sh, f"B{r}")
            self.calc_row += 1

        def heading(text: str) -> None:
            self.put(sh, f"A{self.calc_row}", text, bold=True, color=NAVY)
            self.calc_row += 1

        def domain_raw(d: str, frac: str, units: str, bench: str) -> str:
            return (f's_anchor_share*s_weights_{d}*{frac}+(1-s_anchor_share)*s_weights_{d}'
                    f'*MIN(1,MAX(0,{units}-{bench})/(s_saturation_{d}*c_sat_scale))')

        last_setting = self.settings_last_row
        heading("Inputs and adjustment")
        step("c_changed", "Settings that differ from the shipped policy",
             f"SUMPRODUCT(--(Settings!$B$7:$H${last_setting}<>Defaults!$B$7:$H${last_setting}))")
        step("c_problems", "Rows that need completing",
             'COUNTIF(PUB_problem,"?*")+COUNTIF(GRA_problem,"?*")+COUNTIF(PRE_problem,"?*")')
        step("c_effort", "Research effort (%)", 'IF(f_effort="",s_standard_research_effort,f_effort)')
        step("c_months", "Months active", 'IF(f_months="",12,f_months)')
        step("c_first", "First-year faculty", 'IF(f_first_year="Yes",1,0)')
        step("c_not_evaluated", "Not evaluated (no research assignment)",
             'IF(AND(s_evaluate_zero_assignment="No",OR(c_effort=0,c_months=0)),1,0)')
        step("c_factor", "Adjustment factor",
             "MAX(s_factor_min,MIN(s_factor_max,(c_effort/s_standard_research_effort)*(c_months/12)))")
        step("c_sat_scale", "Scale applied to excellence caps", "MAX(c_factor,0.5)")
        step("c_bcount", "Publication minimum this year", "MAX(1,INT(s_pub_min_count*c_factor*IF(c_first=1,0.5,1)+0.5))")
        step("c_bunits", "Publication benchmark units", "s_benchmark_units_P*c_bcount/s_pub_min_count")
        step("c_pid", "Proposals in development", "N(o_pid)")
        step("c_documented", "Reduced assignment needs documented activity", 'IF(s_reduced_assignment_rule="documented_activity",1,0)')
        step("c_reduced", "Reduced assignment applies", "IF(c_factor<=s_reduced_assignment_factor,1,0)")
        step("c_waived", "Minimums waived outright", "IF(AND(c_documented=0,c_factor<s_reduced_assignment_factor),1,0)")

        heading("Publications")
        step("c_pub_count", "Publications counted toward the minimum", "SUM(PUB_count)")
        step("c_pub_chapters", "Chapter units after the cap",
             'IF(s_chapter_cap="",SUM(PUB_chapter),MIN(SUM(PUB_chapter),s_chapter_cap))')
        step("c_pub_units", "Publication units",
             "SUM(PUB_main)+c_pub_chapters+MIN(SUM(PUB_other),s_other_output_cap)+MIN(SUM(OT_editorial),s_editorial_cap)"
             "+MIN(SUM(PUB_impact),s_impact_cap)")
        step("c_pub_annual", "Minimum met this year", "IF(c_pub_count>=c_bcount,1,0)")
        step("c_pub_need", "Publications needed over three years",
             'IF(s_rolling_benchmark="current_year",3*c_bcount,c_bcount+2*s_pub_min_count)')
        step("c_pub_roll", "Minimum met on the rolling window",
             "IF(AND(c_pub_annual=0,ISNUMBER(o_pubs_prior2),s_rolling_max_domains>0,"
             "c_pub_count+N(o_pubs_prior2)>=c_pub_need,N(o_manuscripts)>=1),1,0)")
        step("c_left1", "Rolling-window rescues left after publications", "s_rolling_max_domains-c_pub_roll")
        step("c_gP", "Publication minimum met", "IF(OR(c_pub_annual=1,c_pub_roll=1),1,0)")
        step("c_fracP", "Share of the publication minimum earned", "IF(c_gP=1,1,MIN(1,c_pub_count/c_bcount))")
        step("c_rawP", "Publication points (unrounded)", domain_raw("P", "c_fracP", "c_pub_units", "c_bunits"))

        heading("Grants and contracts")
        step("c_g_review", "Grant review units after the cap", "MIN(SUM(OT_review),s_grant_review_cap)")
        step("c_g_units", "Grant units",
             "SUM(GRA_awardunits)+MIN(SUM(GRA_subunits),s_submission_cap)"
             "+MIN(s_implementation_units*SUM(GRA_active),s_implementation_cap)+c_g_review"
             "+MIN(s_development_units*(c_pid+N(o_unfunded)),s_development_cap)")
        step("c_g_raw", "Qualifying submission, award or active project",
             "IF(SUM(GRA_meets)+SUM(GRA_awardgate)+SUM(GRA_active)>0,1,0)")
        step("c_g_partial", "Partial grant engagement only",
             "IF(AND(c_g_raw=0,OR(SUM(GRA_subpartial)+SUM(GRA_awardpartial)>0,c_pid>0,c_g_review>0)),1,0)")
        step("c_g_first", "First-year allowance applies", "IF(AND(c_first=1,c_g_raw=0,c_pid>0),1,0)")
        step("c_g4", "Minimum met before the rolling window",
             "IF(OR(c_g_raw=1,c_waived=1,AND(c_reduced=1,c_g_partial=1),c_g_first=1),1,0)")
        step("c_g_roll", "Minimum met on the rolling window",
             "IF(AND(c_g4=0,ISNUMBER(o_grant_prior2),c_left1>0,N(o_grant_prior2)>=2,c_pid>=1),1,0)")
        step("c_left2", "Rolling-window rescues left after grants", "c_left1-c_g_roll")
        step("c_gG", "Grant minimum met", "IF(OR(c_g4=1,c_g_roll=1),1,0)")
        step("c_fracG", "Share of the grant minimum earned", "IF(c_gG=1,1,IF(c_g_partial=1,0.5,0))")
        step("c_rawG", "Grant points (unrounded)", domain_raw("G", "c_fracG", "c_g_units", "s_benchmark_units_G"))
        step("c_distinct", "New large competitive award as PI or MPI", "IF(SUM(GRA_distinct)>0,1,0)")

        heading("Presentations")
        step("c_r_units", "Presentation units",
             "SUM(PRE_units)+MIN(SUM(PRE_localunits),s_local_cap)+MIN(s_chair_units*N(o_chaired),s_chair_cap)")
        step("c_r_raw", "Qualifying conference presentation", "IF(SUM(PRE_gate)>0,1,0)")
        step("c_r_partial", "Partial dissemination activity only", "IF(AND(c_r_raw=0,SUM(PRE_partial)>0),1,0)")
        step("c_r2", "Minimum met before the rolling window",
             "IF(OR(c_r_raw=1,c_waived=1,AND(c_documented=1,c_reduced=1,c_r_partial=1)),1,0)")
        step("c_r_roll", "Minimum met on the rolling window",
             "IF(AND(c_r2=0,ISNUMBER(o_pres_prior2),c_left2>0,N(o_pres_prior2)>=2,N(o_abstracts)>=1),1,0)")
        step("c_gR", "Presentation minimum met", "IF(OR(c_r2=1,c_r_roll=1),1,0)")
        step("c_fracR", "Share of the presentation minimum earned", "IF(c_gR=1,1,IF(c_r_partial=1,0.5,0))")
        step("c_rawR", "Presentation points (unrounded)", domain_raw("R", "c_fracR", "c_r_units", "s_benchmark_units_R"))

        heading("Displayed scores")
        step("c_total", "Overall research score", "ROUND(ROUND(c_rawP+c_rawG+c_rawR,9),1)")
        for d in "PGR":
            step(f"c_own{d}", f"Domain {d} rounded on its own", f"ROUND(ROUND(c_raw{d},9),1)")
        step("c_gap", "Tenths needed for the parts to add up",
             'IF(s_total_from_rounded_parts="Yes",ROUND((c_total-(c_ownP+c_ownG+c_ownR))*10,0),0)')
        for d in "PGR":
            step(f"c_err{d}", f"Domain {d} rounding room in the needed direction", f"(c_raw{d}-c_own{d})*SIGN(c_gap)")
        step("c_showP", "Publications score shown",
             "ROUND(c_ownP+IF(AND(c_gap<>0,c_errP>=c_errG,c_errP>=c_errR),0.1*SIGN(c_gap),0),1)")
        step("c_showG", "Grants score shown",
             "ROUND(c_ownG+IF(AND(c_gap<>0,NOT(AND(c_errP>=c_errG,c_errP>=c_errR)),c_errG>=c_errR),0.1*SIGN(c_gap),0),1)")
        step("c_showR", "Presentations score shown",
             "ROUND(c_ownR+IF(AND(c_gap<>0,NOT(AND(c_errP>=c_errG,c_errP>=c_errR)),NOT(c_errG>=c_errR)),0.1*SIGN(c_gap),0),1)")
        for d in "PGR":
            step(f"c_share{d}", f"Domain {d} share of its points", f"IF(s_weights_{d}=0,0,c_raw{d}/s_weights_{d})")

        heading("Classification")
        step("c_unmet", "Minimums not met", "3-c_gP-c_gG-c_gR")
        step("c_by_total", "Score reaches the Exceeds threshold", "IF(c_total>=s_exceeds_threshold,1,0)")
        step("c_broad", "Two strong domains",
             "IF((c_shareP>=s_breadth_level)+(c_shareG>=s_breadth_level)+(c_shareR>=s_breadth_level)>=2,1,0)")
        step("c_needs_breadth", "Rule needs two strong domains",
             'IF(OR(s_exceeds_rule="breadth",s_exceeds_rule="breadth_or_distinction"),1,0)')
        step("c_on_score", "Exceeds on score", "IF(AND(c_by_total=1,OR(c_broad=1,c_needs_breadth=0)),1,0)")
        step("c_on_award", "Exceeds through a large award",
             'IF(AND(c_distinct=1,c_total>=s_distinction_min_total,OR(s_exceeds_rule="total_or_distinction",'
             's_exceeds_rule="breadth_or_distinction")),1,0)')
        step("c_empty", "Nothing has been entered yet",
             'IF(AND(f_name&f_id="",SUM(PUB_used)+SUM(GRA_used)+SUM(PRE_used)+SUM(OT_editorial)+SUM(OT_review)=0,'
             'N(o_chaired)+N(o_pid)+N(o_unfunded)+N(o_manuscripts)+N(o_abstracts)=0,'
             'COUNT(o_pubs_prior2,o_grant_prior2,o_pres_prior2)=0),1,0)')
        step("c_basis", "Basis for the classification",
             'IF(c_problems>0,"incomplete",IF(c_empty=1,"empty",IF(c_not_evaluated=1,"not_evaluated",IF(c_unmet>0,"minimum_unmet",'
             'IF(AND(s_enforce_meets_threshold="Yes",c_total<s_meets_threshold),"below_meets_threshold",'
             'IF(c_on_score=1,"exceeds_on_score",IF(c_on_award=1,"exceeds_on_award",'
             'IF(AND(c_by_total=1,c_needs_breadth=1),"meets_lacks_breadth","meets"))))))))')
        step("c_class", "Research evaluation",
             'IF(c_basis="incomplete","Incomplete: complete the rows marked on the entry sheets",'
             'IF(c_basis="empty","Nothing entered yet",IF(c_basis="not_evaluated","Not Evaluated",IF(OR(c_basis="minimum_unmet",c_basis="below_meets_threshold"),'
             '"Does Not Meet Expectations",IF(LEFT(c_basis,7)="exceeds","Exceeds Expectations","Meets Expectations")))))')
        step("c_scored", "A score is available", 'IF(OR(c_basis="incomplete",c_basis="empty",c_basis="not_evaluated"),0,1)')
        step("c_approaching", "Approaching expectations",
             'IF(AND(c_basis="minimum_unmet",c_unmet=1,c_total>=s_approaching_threshold),1,0)')
        for d, frac in (("P", "c_fracP"), ("G", "c_fracG"), ("R", "c_fracR")):
            step(f"c_status{d}", f"Domain {d} performance",
                 f'IF(c_g{d}=1,IF(c_share{d}>=s_domain_exceeds_level,"Exceeds","Meets"),IF({frac}>0,"Approaching","Not met"))')

        heading("Wording")
        step("c_unmet_pg", "Unmet minimums (publications, grants)",
             'IF(c_gP=0,"publications","")&IF(AND(c_gP=0,c_gG=0)," and ","")&IF(c_gG=0,"grants and contracts","")')
        step("c_unmet_names", "Unmet minimums",
             'c_unmet_pg&IF(AND(c_unmet_pg<>"",c_gR=0)," and ","")&IF(c_gR=0,"presentations","")')
        step("c_reason", "Why this classification",
             'IF(c_basis="empty","Enter the publications, grants and presentations for the year on the entry sheets.",'
             'IF(c_basis="incomplete","Some rows on the entry sheets are incomplete; see the \'Check this row\' column.",'
             'IF(c_basis="not_evaluated","Not evaluated on research: no assigned research effort or no active months. Refer to the committee.",'
             'IF(c_basis="minimum_unmet","Minimum expectation not met in "&c_unmet_names&"."&IF(c_approaching=1,'
             '" Only one minimum is unmet and the score is at least "&s_approaching_threshold&", so this is flagged for committee review.",""),'
             '"All three minimum expectations are met. "&IF(c_basis="below_meets_threshold","The score of "&TEXT(c_total,"0.0")'
             '&" is still below the Meets threshold of "&s_meets_threshold&".",IF(c_basis="exceeds_on_score",'
             '"The score is at or above the Exceeds threshold of "&s_exceeds_threshold&".",IF(c_basis="exceeds_on_award",'
             '"Exceeds Expectations is awarded for a new competitive large award as PI or MPI; the score itself is below the Exceeds threshold of "'
             '&s_exceeds_threshold&".",IF(c_basis="meets_lacks_breadth","The score reaches the Exceeds threshold of "&s_exceeds_threshold'
             '&", but Exceeds also needs two domains at "&TEXT(s_breadth_level,"0%")&" of their points.",'
             '"Exceeds Expectations starts at "&s_exceeds_threshold&", "&TEXT(s_exceeds_threshold-c_total,"0.0")&" points away."))))))))')
        step("c_basisP", "Basis for the publication minimum",
             'IF(c_pub_annual=1,c_pub_count&" counted toward a minimum of "&c_bcount,IF(c_pub_roll=1,'
             '"three-year rolling window ("&(c_pub_count+N(o_pubs_prior2))&" of "&c_pub_need&" required)",'
             '"not met: "&c_pub_count&" counted toward a minimum of "&c_bcount))')
        step("c_basisG", "Basis for the grant minimum",
             'IF(c_g_roll=1,"three-year rolling window",IF(c_g_first=1,"first-year faculty: documented proposal development",'
             'IF(c_g_raw=1,"qualifying submission, award or active funded project",IF(c_waived=1,"waived: reduced research assignment",'
             'IF(AND(c_reduced=1,c_g_partial=1),"reduced assignment: documented grant engagement","not met")))))')
        step("c_basisR", "Basis for the presentation minimum",
             'IF(c_r_raw=1,"qualifying conference presentation",IF(c_waived=1,"waived: reduced research assignment",'
             'IF(AND(c_documented=1,c_reduced=1,c_r_partial=1),"reduced assignment: documented dissemination activity",'
             'IF(c_r_roll=1,"three-year rolling window","not met"))))')
        step("c_strong_pg", "Strong domains (publications, grants)",
             'IF(c_statusP="Exceeds","Publications","")&IF(AND(c_statusP="Exceeds",c_statusG="Exceeds"),", ","")'
             '&IF(c_statusG="Exceeds","Grants and contracts","")')
        step("c_strengths", "Areas of strength",
             'IF(c_scored=0,"",IF(c_strong_pg&IF(c_statusR="Exceeds","x","")="","No domain above "&TEXT(s_domain_exceeds_level,"0%")'
             '&" of its points this year.",c_strong_pg&IF(AND(c_strong_pg<>"",c_statusR="Exceeds"),", ","")'
             '&IF(c_statusR="Exceeds","Presentations","")))')
        step("c_needs", "Areas requiring improvement",
             'IF(c_scored=0,"",IF(c_unmet=0,"All minimum expectations are met.",UPPER(LEFT(c_unmet_names,1))&MID(c_unmet_names,2,200)))')
        notes = [
            ("c_note1", 'IF(c_scored=0,"",IF(c_first=1,"First-year faculty: publication minimum "&c_bcount&".",IF(c_factor<1,'
                        '"Expectations prorated (adjustment factor "&TEXT(c_factor,"0.00")&"; publication minimum "&c_bcount&").","")))'),
            ("c_note2", 'IF(AND(c_scored=1,c_factor>1),"Expectations scaled up for a research-intensive assignment (factor "'
                        '&TEXT(c_factor,"0.00")&"; publication minimum "&c_bcount&").","")'),
            ("c_note3", 'IF(AND(c_scored=1,c_pub_roll=1),"Publication minimum met on the three-year rolling window, with a manuscript under review.","")'),
            ("c_note4", 'IF(AND(c_scored=1,c_gG=1,c_g_raw=0),"Grant minimum met by allowance: "&c_basisG&".","")'),
            ("c_note5", 'IF(AND(c_scored=1,c_gR=1,c_r_raw=0),"Presentation minimum met by allowance: "&c_basisR&".","")'),
            ("c_flag1", 'IF(AND(c_scored=1,c_unmet>0,c_total>=s_meets_threshold),"Committee review: the score is at or above "'
                        '&s_meets_threshold&" but a required minimum is unmet.",IF(AND(c_scored=1,c_unmet>0,'
                        'MAX(c_shareP,c_shareG,c_shareR)>=s_distinguished_level),"Committee review: distinguished performance in one '
                        'domain with a required minimum unmet.",IF(c_approaching=1,"Committee review: one required minimum is unmet and '
                        'the score is at or above "&s_approaching_threshold&" (approaching expectations).","")))'),
            ("c_flag2", 'IF(AND(c_scored=1,c_basis="below_meets_threshold"),"All minimums are met but the score is below the Meets threshold.","")'),
            ("c_flag3", 'IF(AND(c_scored=1,SUM(GRA_subunits)>s_submission_cap),"Submission credit capped at "&s_submission_cap&" units.","")'),
            ("c_flag4", 'IF(AND(c_scored=1,COUNTIF(PUB_kind,"tier0")>=3),"Committee review: "&COUNTIF(PUB_kind,"tier0")'
                        '&" articles in questionable venues (not credited).","")'),
            ("c_flag5", 'IF(c_scored=0,"",IF(COUNTIF(PUB_note,"?*")+COUNTIF(GRA_note,"?*")+COUNTIF(PRE_note,"?*")>0,'
                        '(COUNTIF(PUB_note,"?*")+COUNTIF(GRA_note,"?*")+COUNTIF(PRE_note,"?*"))&" entry row(s) carry a note; see the Note column on the entry sheets.",""))'),
        ]
        for name, formula in notes:
            step(name, "Explanation line" if name.startswith("c_note") else "Committee line", formula)
        step("c_policy", "Policy statement",
             's_policy_name&" | "&s_policy_note&" | "&IF(c_changed=0,"settings match the shipped policy (fingerprint "&d_fingerprint&")",'
             'c_changed&" setting(s) differ from the shipped policy")')
        self.protect(sh)

    # ------------------------------------------------------------------ dashboard, report, instructions
    def status_rules(self, sheet: str, ref: str, cell: str) -> None:
        ws = self.sheets[sheet]
        good, warn, bad = (PatternFill("solid", bgColor=c) for c in ("CFE8DA", "F6E3B5", "F2CFCB"))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'OR(LEFT({cell},5)="Meets",LEFT({cell},7)="Exceeds")'], fill=good))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'OR({cell}="Approaching",LEFT({cell},3)="Not",LEFT({cell},10)="Incomplete",AND(LEFT({cell},4)="Does",c_approaching=1))'], fill=warn))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'OR({cell}="Not met",AND(LEFT({cell},4)="Does",c_approaching=0))'], fill=bad))

    def results_block(self, sheet: str, top: int) -> int:
        """Write the result summary starting at row `top`; returns the next free row."""
        ws = self.sheets[sheet]
        r = top
        self.put(sheet, f"A{r}", "Research evaluation", bold=True)
        ws.merge_cells(f"B{r}:E{r}")
        self.put(sheet, f"B{r}", "=c_class", bold=True, size=14, border=True)
        self.status_rules(sheet, f"B{r}", f"$B${r}")
        self.put(sheet, f"A{r + 1}", "Overall research score (of 100)", bold=True)
        self.put(sheet, f"B{r + 1}", '=IF(c_scored=1,c_total,"")', bold=True, size=14, fmt="0.0", align="left")
        ws.merge_cells(f"A{r + 2}:E{r + 2}")
        self.put(sheet, f"A{r + 2}", "=c_reason", wrap=True)
        ws.row_dimensions[r + 2].height = 46
        r += 4
        for i, header in enumerate(("Research domain", "Points earned", "Maximum points", "Performance", "Minimum expectation")):
            self.put(sheet, f"{col(i + 1)}{r}", header, bold=True, color="FFFFFF", fill=HEAD_FILL, border=True, wrap=True)
        for i, (d, label) in enumerate((("P", "Publications"), ("G", "Grants and contracts"), ("R", "Presentations")), 1):
            self.put(sheet, f"A{r + i}", label, border=True)
            self.put(sheet, f"B{r + i}", f'=IF(c_scored=1,c_show{d},"")', border=True, fmt="0.0")
            self.put(sheet, f"C{r + i}", f"=s_weights_{d}", border=True, fmt="0.0")
            self.put(sheet, f"D{r + i}", f'=IF(c_scored=1,c_status{d},"")', border=True)
            self.put(sheet, f"E{r + i}", f'=IF(c_scored=1,IF(c_g{d}=1,"Met: ","")&c_basis{d},"")', border=True, wrap=True)
            self.status_rules(sheet, f"D{r + i}", f"$D${r + i}")
            ws.row_dimensions[r + i].height = 30
        self.put(sheet, f"A{r + 4}", "Overall research score", bold=True, border=True)
        self.put(sheet, f"B{r + 4}", '=IF(c_scored=1,c_total,"")', bold=True, border=True, fmt="0.0")
        self.put(sheet, f"C{r + 4}", 100, bold=True, border=True, fmt="0.0")
        self.put(sheet, f"D{r + 4}", "=c_class", bold=True, border=True, wrap=True)
        self.put(sheet, f"E{r + 4}", "", border=True)
        r += 6
        self.put(sheet, f"A{r}", "Areas of strength", bold=True)
        ws.merge_cells(f"B{r}:E{r}")
        self.put(sheet, f"B{r}", "=c_strengths", wrap=True)
        self.put(sheet, f"A{r + 1}", "Areas requiring improvement", bold=True)
        ws.merge_cells(f"B{r + 1}:E{r + 1}")
        self.put(sheet, f"B{r + 1}", "=c_needs", wrap=True)
        r += 3
        self.put(sheet, f"A{r}", "How the outcome was reached", bold=True)
        for i in range(1, 6):
            ws.merge_cells(f"A{r + i}:E{r + i}")
            line = f"=c_note{i}" if i > 1 else ('=IF(AND(c_scored=1,c_note1&c_note2&c_note3&c_note4&c_note5=""),'
                                                '"No allowance or proration was applied.",c_note1)')
            self.put(sheet, f"A{r + i}", line, wrap=True)
        r += 7
        self.put(sheet, f"A{r}", "For committee attention", bold=True)
        for i in range(1, 6):
            ws.merge_cells(f"A{r + i}:E{r + i}")
            line = f"=c_flag{i}" if i > 1 else ('=IF(AND(c_scored=1,c_flag1&c_flag2&c_flag3&c_flag4&c_flag5=""),'
                                                '"Nothing flagged.",c_flag1)')
            self.put(sheet, f"A{r + i}", line, wrap=True)
        return r + 7

    def build_dashboard(self) -> None:
        sh = "Dashboard"
        ws = self.sheets[sh]
        self.title(sh, "Evaluation results", "Updates automatically as data are entered. Nothing on this sheet is typed.")
        self.put(sh, "A3", '=IF(f_name&f_id="","Unnamed faculty member",f_name&IF(AND(f_name<>"",f_id<>""),", ","")&f_id)'
                           '&IF(f_year<>"",", "&f_year,"")', bold=True, size=12)
        r = self.results_block(sh, 5)
        self.put(sh, f"A{r}", "How the points were earned", bold=True, color=NAVY, size=12)
        heads = ("Research domain", "Units earned", "Benchmark units", "Units for full excellence", "Points at the minimum",
                 "Excellence points", "Share of maximum")
        for i, header in enumerate(heads):
            self.put(sh, f"{col(i + 1)}{r + 1}", header, bold=True, fill=CALC_FILL, border=True, wrap=True)
        rows = (("P", "Publications", "c_pub_units", "c_bunits", "c_fracP"), ("G", "Grants and contracts", "c_g_units", "s_benchmark_units_G", "c_fracG"),
                ("R", "Presentations", "c_r_units", "s_benchmark_units_R", "c_fracR"))
        for i, (d, label, units, bench, frac) in enumerate(rows, 2):
            anchor = f"s_anchor_share*s_weights_{d}*{frac}"
            self.put(sh, f"A{r + i}", label, border=True)
            self.put(sh, f"B{r + i}", f'=IF(c_scored=1,{units},"")', border=True, fmt="0.00")
            self.put(sh, f"C{r + i}", f'=IF(c_scored=1,{bench},"")', border=True, fmt="0.00")
            self.put(sh, f"D{r + i}", f'=IF(c_scored=1,s_saturation_{d}*c_sat_scale,"")', border=True, fmt="0.00")
            self.put(sh, f"E{r + i}", f'=IF(c_scored=1,{anchor},"")', border=True, fmt="0.00")
            self.put(sh, f"F{r + i}", f'=IF(c_scored=1,c_raw{d}-{anchor},"")', border=True, fmt="0.00")
            self.put(sh, f"G{r + i}", f"=IF(c_scored=1,c_share{d},0)", border=True, fmt="0%")
        ws.conditional_formatting.add(f"G{r + 2}:G{r + 4}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                      end_value=1, color=NAVY))
        self.put(sh, f"A{r + 6}", '="Meeting a domain\'s minimum earns "&TEXT(s_anchor_share,"0%")&" of its points. Units above the '
                                  'benchmark earn the rest, in proportion, up to the units needed for full excellence. Adjustment factor: "'
                                  '&TEXT(c_factor,"0.00")&"."', wrap=True, color="5A6A79")
        ws.merge_cells(f"A{r + 6}:G{r + 6}")
        ws.row_dimensions[r + 6].height = 30
        self.put(sh, f"A{r + 8}", "Scoring policy", bold=True)
        ws.merge_cells(f"B{r + 8}:G{r + 8}")
        self.put(sh, f"B{r + 8}", "=c_policy", wrap=True)
        ws.merge_cells(f"A{r + 9}:G{r + 9}")
        self.put(sh, f"A{r + 9}", "Decision aid only. The scoring policy is provisional until adopted by the annual review committee; "
                                  "the chair's written evaluation and the faculty narrative remain the record.", wrap=True, color="5A6A79")
        ws.row_dimensions[r + 9].height = 30
        for letter, width in zip("ABCDEFG", (34, 16, 16, 22, 48, 16, 16)):
            ws.column_dimensions[letter].width = width
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.protect(sh)

    def build_report(self) -> None:
        sh = "Report"
        ws = self.sheets[sh]
        self.put(sh, "A1", "Faculty Research Evaluation Report", bold=True, size=16, color=NAVY)
        info = [("Faculty member", "f_name"), ("Faculty ID", "f_id"), ("Academic rank", "f_rank"),
                ("Department", "f_department"), ("Appointment type", "f_appointment"), ("Evaluation period", "f_year"),
                ("Assigned research effort (%)", "c_effort"), ("Months active", "c_months")]
        for i, (label, name) in enumerate(info):
            self.put(sh, f"A{3 + i}", label, bold=True)
            ws.merge_cells(f"B{3 + i}:E{3 + i}")
            self.put(sh, f"B{3 + i}", f'=IF({name}="","",{name})', align="left")
        r = self.results_block(sh, 12)
        activity = [("Publication productivity", '=c_pub_count&" counted toward the minimum; "&TEXT(c_pub_units,"0.00")&" units"'),
                    ("Grant activity and awards", '=SUM(GRA_issub)&" proposal(s) submitted, "&SUM(GRA_isaward)&" award(s) received, "'
                                                  '&SUM(GRA_active)&" active project(s); "&TEXT(c_g_units,"0.00")&" units"'),
                    ("Scientific presentations", '=SUM(PRE_ok)&" presentation(s) entered; "&TEXT(c_r_units,"0.00")&" units"'),
                    ("Additional scholarly contributions", '=TEXT(MIN(SUM(OT_editorial),s_editorial_cap),"0.00")&" editorial units, "'
                                                           '&TEXT(c_g_review,"0.00")&" grant-review units, "&N(o_chaired)&" session(s) chaired"')]
        self.put(sh, f"A{r}", "Activity summary", bold=True, color=NAVY, size=12)
        for i, (label, formula) in enumerate(activity, 1):
            self.put(sh, f"A{r + i}", label, bold=True)
            ws.merge_cells(f"B{r + i}:E{r + i}")
            self.put(sh, f"B{r + i}", formula, wrap=True)
        r += len(activity) + 2
        self.put(sh, f"A{r}", "Recorded for the committee (not scored)", bold=True, color=NAVY, size=12)
        for i, (label, name) in enumerate((("Awards and recognition", "x_awards"), ("Mentoring", "x_mentoring"),
                                           ("Community-engaged work", "x_community"),
                                           ("Evaluator notes", "x_notes")), 1):
            self.put(sh, f"A{r + i}", label, bold=True)
            ws.merge_cells(f"B{r + i}:E{r + i}")
            self.put(sh, f"B{r + i}", f'=IF({name}="","",{name})', wrap=True)
            ws.row_dimensions[r + i].height = 34
        r += 6
        ws.merge_cells(f"A{r}:E{r}")
        self.put(sh, f"A{r}", '="Scoring policy: "&c_policy', wrap=True, color="5A6A79", size=9)
        ws.row_dimensions[r].height = 26
        ws.merge_cells(f"A{r + 1}:E{r + 1}")
        self.put(sh, f"A{r + 1}", "Decision aid only. The chair's written evaluation and the faculty narrative remain the record.",
                 color="5A6A79", size=9)
        self.put(sh, f"A{r + 4}", "Evaluator signature: ______________________________")
        self.put(sh, f"D{r + 4}", "Date: ______________")
        for letter, width in zip("ABCDE", (34, 14, 14, 20, 46)):
            ws.column_dimensions[letter].width = width
        ws.print_area = f"A1:E{r + 4}"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        self.protect(sh)

    def build_instructions(self) -> None:
        sh = "Instructions"
        self.title(sh, "Faculty Research Evaluation Calculator",
                   "A decision aid for the annual review of research and scholarship. It informs the chair's written evaluation; "
                   "it does not replace it.")
        lines = [
            ("How to use it", None),
            ("1.", "Faculty sheet: enter the faculty member, the evaluation year, the assigned research effort and the months active."),
            ("2.", "Publications, Grants and Presentations sheets: one row per item. Pick from the dropdown lists in the yellow cells."),
            ("3.", "Other sheet: editorial service, grant review panels, counts for the year and, if needed, the three-year rolling window."),
            ("4.", "Dashboard sheet: the score, the classification and how each point was earned. It updates as you type."),
            ("5.", "Report sheet: a one-page summary. Use File > Print, or File > Save As > PDF."),
            ("", ""),
            ("Reading the colours", None),
            ("Yellow", "Cells you fill in."),
            ("Grey", "Calculated cells. They are locked."),
            ("Green / amber / red", "On the Dashboard: a minimum met or exceeded; approaching; not met."),
            ("", ""),
            ("If a row is not counted", None),
            ("", "Look at the grey 'Check this row' column beside it. It says what is missing, and the Dashboard shows 'Incomplete' until every started row is finished."),
            ("", ""),
            ("How the score works", None),
            ("", "Three domains: publications, grants and contracts, and presentations. Each has a minimum expectation. Meeting it earns a fixed share of the domain's points; work above the minimum earns the rest."),
            ("", "Any unmet minimum gives Does Not Meet Expectations. With every minimum met the result is Meets Expectations, or Exceeds Expectations under the rule on the Settings sheet."),
            ("", ""),
            ("Journal tiers", None),
            ("Tier 1", "Top quartile (Q1) of its subject category, or the flagship journal of a major society in the field."),
            ("Tier 2", "Second quartile (Q2)."),
            ("Tier 3", "Any other peer-reviewed journal indexed in MEDLINE, Scopus or Web of Science."),
            ("Questionable", "No verifiable peer review. Earns nothing and does not count toward the minimum."),
            ("", ""),
            ("Venue tiers", None),
            ("Tier A", "Major national or international society meeting with competitive peer-reviewed abstract selection."),
            ("Tier B", "Regional or state society meeting, or a national meeting with light review."),
            ("Tier C", "Local meeting, or limited review."),
            ("", ""),
            ("The scoring policy", None),
            ("", "Every threshold and criterion is on the Settings sheet. It belongs to the annual review committee and is protected. The values shipped are provisional and are documented in PROVISIONAL_POLICY.md."),
            ("", "The Dashboard states whether the settings still match the shipped policy."),
            ("", ""),
            ("Limits", None),
            ("", f"Up to {ROWS['pubs']} publications, {ROWS['grants']} grants and {ROWS['pres']} presentations per faculty member."),
            ("", "One workbook holds one faculty member for one year. Save a copy per person."),
            ("", "Scores round half up. The original Phase 1 rounding rule is available only in the web application."),
        ]
        ws = self.sheets[sh]
        for i, (lead, text) in enumerate(lines, 4):
            if text is None:
                self.put(sh, f"A{i}", lead, bold=True, color=NAVY, size=12)
            else:
                self.put(sh, f"A{i}", lead, bold=True)
                self.put(sh, f"B{i}", text, wrap=True)
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 110
        self.protect(sh)

    def build(self) -> Workbook:
        self.build_lists()
        self.build_settings()
        self.build_faculty_and_other()
        self.build_publications()
        self.build_grants()
        self.build_presentations()
        self.build_calc()
        self.build_dashboard()
        self.build_report()
        self.build_instructions()
        self.protect("Settings")
        self.protect("Defaults")
        self.wb.active = 0
        return self.wb


def main(path: Path = OUTPUT, policy: dict = PROVISIONAL) -> Path:
    Calculator(policy).build().save(path)
    return path


if __name__ == "__main__":
    print(f"Wrote {main(Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT)}")
