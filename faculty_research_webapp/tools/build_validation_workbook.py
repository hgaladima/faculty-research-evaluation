"""Build faculty_evaluation_validation.xlsx: Python engine vs Excel calculator vs web application.

Run from faculty_research_webapp/:  python -m tools.build_validation_workbook [random cases per policy]

Needs Excel (the calculator is recalculated through it). The web column is read from the running app's result
panel, driven by Streamlit's test harness, with storage redirected to a temporary folder.
"""
from __future__ import annotations

import re
import sys
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from streamlit.testing.v1 import AppTest

import storage
from engine import PROVISIONAL, evaluate, fingerprint
from engine.scenarios import SCENARIOS
from tools import validate_excel
from tools.excel_com import recalculate_and_save

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "faculty_evaluation_validation.xlsx"
APP = str(Path(__file__).resolve().parents[1] / "app.py")
HEAD = PatternFill("solid", fgColor="1D4E89")
FONT = "Arial"


def web_results() -> dict:
    """Load each of the 22 profiles in the web app and read what its result panel displays."""
    tmp = Path(tempfile.mkdtemp(prefix="fre_web_"))
    storage.DATA_DIR, storage.POLICY_FILE = tmp, tmp / "policy.json"
    storage.LOG_FILE, storage.EVAL_DIR = tmp / "policy_log.jsonl", tmp / "evaluations"
    out = {}
    at = AppTest.from_file(APP, default_timeout=120).run()
    for scenario in SCENARIOS:
        at.sidebar.selectbox[1].select(scenario["name"]).run()
        next(b for b in at.button if b.label == "Load example profile").click().run()
        panel = next(h.proto.body for h in at.get("html") if "fre-result" in h.proto.body)
        points = re.findall(r'<div class="pts">([\d.]+) <span>', panel)
        out[scenario["name"]] = {"class": re.search(r'<p class="class">([^<]+)</p>', panel).group(1),
                                 "total": float(re.search(r'data-score="([\d.]+)"', panel).group(1)),
                                 "P": float(points[0]), "G": float(points[1]), "R": float(points[2])}
    return out


def write(ws, row: int, values: list, bold=False, head=False, wrap=False) -> None:
    for i, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=value)
        cell.font = Font(name=FONT, bold=bold or head, color="FFFFFF" if head else "14202B")
        cell.alignment = Alignment(wrap_text=wrap or head, vertical="top")
        if head:
            cell.fill = HEAD


def build(random_per_policy: int = 150) -> Path:
    excel_runs = validate_excel.run(random_per_policy)
    web = web_results()
    wb = Workbook()

    sc = wb.active
    sc.title = "22 scenarios"
    write(sc, 1, ["Scenario", "Python P", "Python G", "Python R", "Python total", "Python class",
                  "Excel P", "Excel G", "Excel R", "Excel total", "Excel class",
                  "Web P", "Web G", "Web R", "Web total", "Web class",
                  "Largest score difference", "Classification match"], head=True)
    excel_by_case = {r["case"]: r["excel"] for r in excel_runs["Provisional"]}
    for i, scenario in enumerate(SCENARIOS, 2):
        py, xl, wr = evaluate(scenario, PROVISIONAL), excel_by_case[scenario["name"]], web[scenario["name"]]
        write(sc, i, [scenario["name"], py["P"], py["G"], py["R"], py["total"], py["class"],
                      xl["P"], xl["G"], xl["R"], xl["total"], xl["class"],
                      wr["P"], wr["G"], wr["R"], wr["total"], wr["class"],
                      f"=MAX(ABS(B{i}-G{i}),ABS(C{i}-H{i}),ABS(D{i}-I{i}),ABS(E{i}-J{i}),"
                      f"ABS(B{i}-L{i}),ABS(C{i}-M{i}),ABS(D{i}-N{i}),ABS(E{i}-O{i}))",
                      f'=IF(AND(F{i}=K{i},F{i}=P{i}),"Yes","No")'])
    last_sc = len(SCENARIOS) + 1
    for letter, width in zip("ABCDEFGHIJKLMNOPQR", (44, 9, 9, 9, 11, 28, 9, 9, 9, 11, 28, 9, 9, 9, 11, 28, 14, 14)):
        sc.column_dimensions[letter].width = width
    sc.freeze_panes = "B2"

    rc = wb.create_sheet("All Excel cases")
    write(rc, 1, ["Policy", "Case", "Outcome type", "Python P", "Python G", "Python R", "Python total", "Python class",
                  "Excel P", "Excel G", "Excel R", "Excel total", "Excel class", "Minimums agree",
                  "Largest score difference", "Match"], head=True)
    row = 2
    for policy_name, records in excel_runs.items():
        for r in records:
            py, xl = r["python"], r["excel"]
            kind = "Scored" if py["scored"] else py["class"]
            numbers = lambda d: [d.get(k) for k in ("P", "G", "R", "total")] if d["scored"] else [None] * 4
            same_minimums = py.get("gates") == xl.get("gates") and py.get("approaching") == xl.get("approaching")
            gates = "Yes" if same_minimums or not py["scored"] else "No"
            write(rc, row, [policy_name, r["case"], kind, *numbers(py), py["class"], *numbers(xl), xl["class"], gates,
                            f'=IF(C{row}="Scored",MAX(ABS(D{row}-I{row}),ABS(E{row}-J{row}),ABS(F{row}-K{row}),ABS(G{row}-L{row})),0)',
                            f'=IF(AND(H{row}=M{row},N{row}="Yes",O{row}<0.05),"Yes","No")'])
            row += 1
    last_rc = row - 1
    for letter, width in zip("ABCDEFGHIJKLMNOP", (26, 44, 16, 9, 9, 9, 11, 28, 9, 9, 9, 11, 28, 12, 14, 9)):
        rc.column_dimensions[letter].width = width
    rc.freeze_panes = "C2"

    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = "Validation: Python engine, Excel calculator and web application"
    sm["A1"].font = Font(name=FONT, bold=True, size=15, color="1D4E89")
    write(sm, 2, [f"Run on {date.today():%d %B %Y}. Shipped provisional policy, fingerprint {fingerprint(PROVISIONAL)}."])
    write(sm, 4, ["Check", "Cases", "Mismatches"], head=True)
    write(sm, 5, ["22 scenarios: classification agrees across Python, Excel and the web app", f"=COUNTA('22 scenarios'!A2:A{last_sc})",
                  f"=COUNTIF('22 scenarios'!R2:R{last_sc},\"No\")"])
    write(sm, 6, ["22 scenarios: largest score difference across the three (points)", f"=COUNTA('22 scenarios'!A2:A{last_sc})",
                  f"=MAX('22 scenarios'!Q2:Q{last_sc})"])
    r = 7
    for policy_name in excel_runs:
        write(sm, r, [f"Excel vs Python, policy '{policy_name}': scenarios and random cases",
                      f"=COUNTIF('All Excel cases'!A2:A{last_rc},\"{policy_name}\")",
                      f"=COUNTIFS('All Excel cases'!A2:A{last_rc},\"{policy_name}\",'All Excel cases'!P2:P{last_rc},\"No\")"])
        r += 1
    write(sm, r, ["All Excel vs Python cases", f"=COUNTA('All Excel cases'!B2:B{last_rc})",
                  f"=COUNTIF('All Excel cases'!P2:P{last_rc},\"No\")"], bold=True)
    r += 2
    notes = [
        "What each column is",
        "Python: the scoring engine called directly on the scenario.",
        "Excel: the case typed into faculty_research_calculator.xlsx through Excel itself, recalculated, and read back.",
        "Web: the scenario loaded in the web application; the values are read from its result panel.",
        "",
        "What 'match' means",
        "The three domain scores, the overall score, the classification, which minimums are met and the 'approaching' label all agree.",
        "",
        "The policies",
        "Provisional is the shipped policy. The other three flip every rule switch the workbook implements, so that the Settings sheet is "
        "shown to drive the formulas: Phase 1 rules with half-up rounding, a mixed set of switches, and a total-only Exceeds rule.",
        "",
        "Random cases",
        "Each is a randomly filled entry form, including incomplete rows, unverified entries, duplicate DOIs, prior-year marks, "
        "reduced assignments, first-year faculty and the rolling window. 'Incomplete' and 'Not Evaluated' cases must be reported as such by both.",
        "",
        "Limits",
        "Excel always rounds half up. The original Phase 1 rule (round half to even) exists only in the Python engine and the web app, "
        "so a score ending in exactly 5 at the second decimal can differ by 0.1 from the unmodified Phase 1 preset.",
        "The random cases are generated, not real faculty records. A back-test on past reviews is still outstanding.",
        "The web application was checked on the 22 scenarios only.",
    ]
    for text in notes:
        heading = text in ("What each column is", "What 'match' means", "The policies", "Random cases", "Limits")
        write(sm, r, [text], bold=heading, wrap=True)
        r += 1
    sm.column_dimensions["A"].width = 110
    sm.column_dimensions["B"].width = 10
    sm.column_dimensions["C"].width = 12
    wb.save(OUTPUT)
    errors = recalculate_and_save(OUTPUT)
    if errors:
        raise AssertionError(f"Formula errors in the validation workbook: {errors[:10]}")
    shown = [row[15] for row in load_workbook(OUTPUT, data_only=True)["All Excel cases"].iter_rows(min_row=2, values_only=True)]
    harness = ["Yes" if r["match"] else "No" for records in excel_runs.values() for r in records]
    if shown != harness:
        raise AssertionError("The Match column of the workbook disagrees with the comparison that produced it.")
    return OUTPUT


if __name__ == "__main__":
    path = build(int(sys.argv[1]) if len(sys.argv) > 1 else 150)
    print(f"Wrote {path}")
