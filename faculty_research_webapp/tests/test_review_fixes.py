"""One or more tests per finding of the application-layer review (numbered as reported)."""
import copy
import io
import tomllib
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

import exports
import forms
import storage
from engine import PROVISIONAL, check_policy, evaluate
from engine.research_eval import NOT_EVALUATED
from engine.scenarios import SCENARIOS, A
from test_app import APP

APP_DIR = Path(APP).parent
META = {"version": 1, "locked": False, "approved_by": ""}


@pytest.fixture
def data_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    monkeypatch.setattr(storage, "POLICY_FILE", tmp_path / "policy.json")
    monkeypatch.setattr(storage, "LOG_FILE", tmp_path / "policy_log.jsonl")
    monkeypatch.setattr(storage, "EVAL_DIR", tmp_path / "evaluations")
    return tmp_path


def run_app():
    return AppTest.from_file(APP, default_timeout=60).run()


def button(at, label):
    return next(b for b in at.button if b.label == label)


def changed(**settings):
    return dict(copy.deepcopy(PROVISIONAL), **settings)


def build(pubs=(), grants=(), pres=()):
    return forms.build_faculty({}, list(pubs), list(grants), list(pres), [], [], PROVISIONAL)


# 1 ---------------------------------------------------------------------------------------------
def test_1_app_listens_on_this_machine_only():
    config = tomllib.loads((APP_DIR / ".streamlit" / "config.toml").read_text(encoding="utf-8"))
    assert config["server"]["address"] == "localhost"


# 2 ---------------------------------------------------------------------------------------------
@pytest.mark.parametrize("path, value", [
    (("journal_tier_mult", 1), float("nan")), (("grant_role_mult", "PI"), None), (("weights", "P"), float("nan")),
    (("meets_threshold",), None), (("saturation", "G"), float("inf")), (("exceeds_rule",), ""),
    (("enforce_meets_threshold",), "yes"),
])
def test_2_blank_or_non_numeric_policy_values_are_rejected(path, value, data_dir):
    policy = copy.deepcopy(PROVISIONAL)
    node = policy
    for part in path[:-1]:
        node = node[part]
    node[path[-1]] = value
    errors, _ = check_policy(policy)
    assert errors and "blank or not valid" in errors[0]
    with pytest.raises(ValueError):
        storage.save_policy(storage.load_policy_state(), policy, who="Chair", year="2026", note="", lock=True)
    assert not storage.POLICY_FILE.exists()


def test_2_an_empty_chapter_cap_is_still_allowed():
    assert check_policy(changed(chapter_cap=None))[0] == []


# 3 ---------------------------------------------------------------------------------------------
def test_3_stale_session_cannot_overwrite_a_locked_policy(data_dir):
    stale = storage.load_policy_state()
    locked = storage.save_policy(stale, changed(exceeds_threshold=90.0), who="Dean", year="2026", note="", lock=True)
    with pytest.raises(storage.PolicyConflict):
        storage.save_policy(stale, changed(meets_threshold=65.0), who="Other", year="2026", note="", lock=False)
    with pytest.raises(storage.PolicyConflict):
        storage.unlock_policy(stale, who="Other", reason="stale")
    on_disk = storage.load_policy_state()
    assert on_disk["meta"]["locked"] and on_disk["policy"]["exceeds_threshold"] == 90.0
    assert storage.state_key(on_disk) == storage.state_key(locked)


def test_3_a_locked_policy_must_be_unlocked_before_saving(data_dir):
    locked = storage.save_policy(storage.load_policy_state(), changed(), who="Dean", year="2026", note="", lock=True)
    with pytest.raises(storage.PolicyConflict):
        storage.save_policy(locked, changed(meets_threshold=65.0), who="Dean", year="2026", note="", lock=False)
    reopened = storage.unlock_policy(locked, who="Dean", reason="correction")
    saved = storage.save_policy(reopened, changed(exceeds_threshold=88.0), who="Dean", year="2026", note="",
                                lock=True)
    assert saved["meta"]["version"] == 2
    assert [e["action"] for e in storage.read_log()] == ["approved and locked", "unlocked", "approved and locked"]


def test_3_an_open_page_picks_up_a_policy_saved_elsewhere(data_dir):
    at = run_app()
    storage.save_policy(storage.load_policy_state(), changed(policy_name="Adopted 2026"), who="Dean", year="2026",
                        note="", lock=True)
    at.run()
    assert not at.exception
    assert any("changed in another window" in w.value for w in at.warning)
    assert any("Adopted 2026" in h.proto.body for h in at.get("html"))


# 4 ---------------------------------------------------------------------------------------------
def test_4_entered_text_is_never_written_as_an_excel_formula():
    result = evaluate(dict(SCENARIOS[0], publications=[A(3, title="=SUM(A1:A9)"), A(3)]), PROVISIONAL)
    info = {"name": "=1+1", "faculty_id": '=HYPERLINK("http://x","click")', "evaluation_year": "2026"}
    data = exports.excel_bytes(info, result, changed(policy_name="=2+2"), META, {"Evaluator notes": "= strong year"})
    cells = [c for ws in load_workbook(io.BytesIO(data)).worksheets for row in ws.iter_rows() for c in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert {"=1+1", "= strong year", "=2+2"} <= {c.value for c in cells}


# 5 ---------------------------------------------------------------------------------------------
@pytest.mark.parametrize("title", ["x" * 6000, "word " * 1500, "pasted\x0btext\x00here"])
def test_5_reports_build_for_awkward_entries(title):
    result = evaluate(dict(SCENARIOS[0], publications=[A(3, title=title), A(3)]), PROVISIONAL)
    info, context = {"name": title[:50]}, {"Evaluator notes": title}
    assert exports.pdf_bytes(info, result, META, context)[:4] == b"%PDF"
    assert exports.excel_bytes(info, result, PROVISIONAL, META, context)[:2] == b"PK"


# 6 ---------------------------------------------------------------------------------------------
@pytest.mark.parametrize("content", ["", "{not json", "[1, 2]", '{"policy": {"weights": {"P": 50}}, "meta": {}}'])
def test_6_unusable_policy_file_falls_back_with_a_message(content, data_dir):
    storage.POLICY_FILE.write_text(content, encoding="utf-8")
    state = storage.load_policy_state()
    assert state["problem"] and state["policy"] == storage.default_policy_state()["policy"]
    at = run_app()
    assert not at.exception
    assert any("could not be used" in e.value for e in at.error)
    assert storage.POLICY_FILE.read_text(encoding="utf-8") == content


def test_6_saving_over_an_unusable_policy_file_keeps_a_copy(data_dir):
    storage.POLICY_FILE.write_text("{not json", encoding="utf-8")
    storage.save_policy(storage.load_policy_state(), changed(), who="Dean", year="2026", note="", lock=False)
    assert storage.load_policy_state()["problem"] is None
    assert len(list(data_dir.glob("policy.unusable-*.json"))) == 1
    assert not list(data_dir.glob("*.tmp"))


def test_6_downloaded_policy_file_can_be_restored(data_dir):
    from engine.policy import to_json
    storage.POLICY_FILE.write_text(to_json(changed(exceeds_threshold=88.0)), encoding="utf-8")
    state = storage.load_policy_state()
    assert state["problem"] is None and state["policy"]["exceeds_threshold"] == 88.0
    assert state["meta"]["locked"] is False


def test_6_damaged_log_and_evaluation_files_do_not_crash(data_dir):
    storage.save_policy(storage.load_policy_state(), changed(), who="Dean", year="2026", note="", lock=False)
    with storage.LOG_FILE.open("a", encoding="utf-8") as f:
        f.write("{broken\n")
    assert len(storage.read_log()) == 1
    storage.EVAL_DIR.mkdir(parents=True)
    (storage.EVAL_DIR / "Bad_2026.json").write_text("{broken", encoding="utf-8")
    (storage.EVAL_DIR / "Odd_2026.json").write_text('{"result": {}}', encoding="utf-8")
    assert storage.load_evaluation("Bad_2026") is None and storage.load_evaluation("Odd_2026") is None
    assert storage.saved_subject("Bad_2026") == "a file that cannot be read"
    at = run_app()
    at.sidebar.selectbox[0].select("Bad_2026").run()
    button(at, "Open saved evaluation").click().run()
    assert not at.exception
    assert any("could not be read" in e.value for e in at.sidebar.error)


# 7 ---------------------------------------------------------------------------------------------
def test_7_only_a_locked_policy_carries_an_approver(data_dir):
    result = evaluate(SCENARIOS[0], PROVISIONAL)
    draft = storage.save_policy(storage.load_policy_state(), changed(exceeds_threshold=88.0), who="Analyst",
                                year="2026", note="", lock=False)
    assert draft["meta"]["approved_by"] == "" and draft["meta"]["saved_by"] == "Analyst"
    assert "not approved" in exports.policy_stamp(result, draft["meta"])
    locked = storage.save_policy(draft, draft["policy"], who="Dean A", year="2026", note="", lock=True)
    assert "locked, approved by Dean A" in exports.policy_stamp(result, locked["meta"])
    reopened = storage.unlock_policy(locked, who="Chair B", reason="typo")
    assert reopened["meta"]["approved_by"] == ""
    assert "Dean A" not in exports.policy_stamp(result, reopened["meta"])


def test_7_policy_cannot_be_saved_without_a_name(data_dir):
    at = run_app()
    assert button(at, "Approve and lock policy").disabled
    assert button(at, "Save as working draft (not locked)").disabled


# 8 ---------------------------------------------------------------------------------------------
GRANT = {"Title": "R21", "Source": "External", "Total amount ($)": 100000.0, "Role": "PI", "Competitive": True}


def test_8_prior_year_tick_on_a_submission_removes_its_credit():
    fac, problems = build(grants=[dict(GRANT, Status=forms.SUBMITTED, **{"Credited in a prior year": True})])
    assert problems == [] and fac["submissions"] == [] and len(fac["entry_notes"]) == 1
    assert evaluate(fac, PROVISIONAL)["gates"]["G"] is False


def test_8_prior_year_award_is_not_recredited_as_a_submission():
    row = dict(GRANT, Status=forms.AWARDED, **{"Also submitted this year": True, "Credited in a prior year": True})
    fac, problems = build(grants=[row])
    assert problems == [] and fac["submissions"] == []
    assert fac["awards"][0]["counted_in_prior_year"] is True
    assert evaluate(fac, PROVISIONAL)["units"]["G"] == 0.0


def test_8_contradictory_ticks_on_an_active_project_are_reported():
    _, problems = build(grants=[dict(GRANT, Status=forms.ACTIVE, **{"Also submitted this year": True})])
    assert len(problems) == 1 and "awarded in a prior year" in problems[0]


# 9 ---------------------------------------------------------------------------------------------
def test_9_not_evaluated_result_can_be_saved_and_exported(data_dir):
    at = run_app()
    at.text_input(key="info_name").set_value("Dr. Admin Leave").run()
    at.number_input(key="info_research_effort").set_value(0.0).run()
    assert any(NOT_EVALUATED in w.value for w in at.warning)
    assert {"Download report (PDF)", "Download report (Excel)"} <= {b.label for b in at.get("download_button")}
    button(at, "Save evaluation").click().run()
    assert not at.exception and storage.list_evaluations() == ["Dr_Admin_Leave_2026"]


def test_9_unscored_results_produce_reports():
    result = evaluate({"research_effort": 0, "name": "Dr. Admin Leave"}, PROVISIONAL)
    assert result["class"] == NOT_EVALUATED
    pdf = exports.pdf_bytes({"name": "Dr. Admin Leave", "research_effort": 0.0}, result, META, {})
    text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert NOT_EVALUATED in text
    assert exports.excel_bytes({"name": "x"}, result, PROVISIONAL, META, {})[:2] == b"PK"


# 10 --------------------------------------------------------------------------------------------
def test_10_names_in_any_script_get_distinct_file_names(data_dir):
    names = [storage.evaluation_name({"name": n, "evaluation_year": "2026"}) for n in ("李明", "Нина Петрова", "José é")]
    assert len(set(names)) == 3 and "unnamed" not in "".join(names)
    assert storage.evaluation_name({"name": "../..\\secret", "evaluation_year": "2026"}) == "secret_2026"
    long_name = storage.save_evaluation({"form": {"info": {"faculty_id": "A" * 300, "evaluation_year": "2026"}}})
    assert len(long_name) <= storage.MAX_NAME_LENGTH + 5 and storage.load_evaluation(long_name) is not None


# 11 --------------------------------------------------------------------------------------------
def test_11_pdf_keeps_zero_values_line_breaks_and_wider_scripts():
    result = evaluate(dict(SCENARIOS[0], research_effort=40), PROVISIONAL)
    info = {"name": "Нина Петрова", "research_effort": 0.0, "months_active": 0}
    pdf = exports.pdf_bytes(info, result, META, {"Evaluator notes": "First line\nSecond line"})
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert "0.0" in text.split("Assigned research effort (%)")[1].split("Months active")[0]
    assert exports.clean(0) == "0" and exports.clean(None) == ""
    assert "First line" in text and "Second line" in text and "First line Second line" not in text
    if exports.pdf_fonts()[0] == "ReportBody":
        assert "Нина Петрова" in text


# 12 --------------------------------------------------------------------------------------------
def test_12_saved_evaluation_appears_in_the_sidebar_at_once(data_dir):
    at = run_app()
    at.text_input(key="info_faculty_id").set_value("F-1").run()
    button(at, "Save evaluation").click().run()
    assert "F_1_2026" in at.sidebar.selectbox[0].options
    assert any("Saved as F_1_2026" in s.value for s in at.success)


# 13 --------------------------------------------------------------------------------------------
@pytest.mark.parametrize("fac", [
    {"publications": [{"type": "book"}, A(3), A(3)]},
    {"publications": [{"type": "article"}, {"type": "article", "role": "lead"}]},
    {"presentations": [{"type": "oral"}]},
    {"presentations": [{"type": "poster", "venue": "A"}]},
    {"awards": [{"size": "major"}], "active_projects": [{}]},
    {"submissions": [{"size": "major", "role": "PI", "id": "X", "scored": True}],
     "awards": [{"size": "major", "role": "PI", "id": "X"}]},
    {"submissions": [{"size": "large", "role": "PI", "id": "X"}],
     "awards": [{"size": "large", "role": "PI", "id": "X", "counted_in_prior_year": True}]},
], ids=["book-no-role", "article-no-tier", "oral-no-venue-or-role", "poster-no-role", "grant-no-role",
        "funded-same-year", "prior-year-award-resubmitted"])
def test_13_incomplete_records_round_trip_through_the_forms(fac):
    f = forms.rows_from_faculty(fac)
    built, problems = forms.build_faculty(f["info"], f["pubs"], f["grants"], f["pres"], f["editorial"],
                                          f["reviews"], PROVISIONAL)
    assert problems == []
    direct, via_forms = evaluate(fac, PROVISIONAL), evaluate(built, PROVISIONAL)
    for key in ("P", "G", "R", "total", "class", "gates", "units"):
        assert via_forms[key] == direct[key], key
    pending = [sum("pending verification" in x for x in r["flags"]) for r in (direct, via_forms)]
    assert pending[0] == pending[1]


# 14 --------------------------------------------------------------------------------------------
def test_14_breakdown_tables_hold_one_type_per_column(data_dir):
    at = run_app()
    at.sidebar.selectbox[1].select("S19 Exceptional across all domains").run()
    button(at, "Load example profile").click().run()
    frames = [d.value for d in at.dataframe if "Counts toward minimum" in d.value.columns]
    assert frames and all(isinstance(v, str) for f in frames for v in f["Counts toward minimum"])
